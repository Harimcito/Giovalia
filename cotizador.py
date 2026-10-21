"""Núcleo del cotizador: lee la cotización en Excel, busca fotos y fichas en la
biblioteca local y arma un PDF único: cotización + fotos + fichas técnicas.

Estructura de la biblioteca (carpeta "biblioteca" por defecto):

    biblioteca/
      indice.json            <- claves, nombres y alias aprendidos
      CF283A/foto.jpg        <- foto del producto (jpg, png o webp)
      CF283A/ficha.pdf       <- ficha técnica en PDF
      TK-3182/foto.png
      ...
"""
import csv
import io
import json
import math
import os
import re
import shlex
import shutil
import subprocess
import tempfile
import unicodedata
from copy import copy
from datetime import date
from pathlib import Path
from xml.sax.saxutils import escape

from openpyxl import load_workbook
from openpyxl.styles import Alignment
from openpyxl.worksheet.properties import PageSetupProperties
from PIL import Image as PILImage
from pypdf import PdfReader, PdfWriter

from buscador import es_codigo, limpiar_consulta, menciones_clave
import pdf_cotizacion
from reportlab.lib import colors
from reportlab.lib.pagesizes import letter
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import cm
from reportlab.platypus import Image, Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

# Líneas de la cotización que no son productos (no llevan foto ni ficha)
PALABRAS_NO_PRODUCTO = ("despacho", "flete", "envio", "transporte", "instalacion", "servicio")
EXT_FOTO = (".jpg", ".jpeg", ".png", ".webp")
# Tokens que parecen un código pero son medidas (64GB, 500ML, 75GR...)
_MEDIDA = re.compile(r"^(\d+(GB|TB|MB|ML|GR|G|KG|MM|CM|W|V|MAH|HZ|MP|PX)|\d+(X\d+)+(CM|MM|M)?)$")


def normalizar(texto) -> str:
    """Mayúsculas, sin tildes ni símbolos: 'TK-3182' -> 'TK3182'."""
    t = unicodedata.normalize("NFKD", str(texto)).encode("ascii", "ignore").decode()
    return re.sub(r"[^A-Z0-9]", "", t.upper())


# ----------------------------------------------------------------------------
# 1. Lectura del Excel
# ----------------------------------------------------------------------------
def leer_cotizacion(ruta) -> list[dict]:
    """Devuelve las líneas de la cotización: cantidad, descripción y precios.

    Ubica la tabla por sus encabezados ('Cantidad', 'Descripción', 'P. Unitario'),
    así que no depende de filas ni columnas fijas.
    """
    ws = load_workbook(ruta).active
    fila_enc, cols = None, {}
    for fila in ws.iter_rows():
        etiquetas = {normalizar(c.value): c.column for c in fila if isinstance(c.value, str)}
        if "CANTIDAD" in etiquetas and "DESCRIPCION" in etiquetas:
            fila_enc, cols = fila[0].row, etiquetas
            break
    if fila_enc is None:
        raise ValueError("No encontré los encabezados 'Cantidad' y 'Descripción' en la hoja.")

    c_cant, c_desc, c_unit = cols["CANTIDAD"], cols["DESCRIPCION"], cols.get("PUNITARIO")
    lineas = []
    for r in range(fila_enc + 1, ws.max_row + 1):
        cant, desc = ws.cell(r, c_cant).value, ws.cell(r, c_desc).value
        if not isinstance(cant, (int, float)) or cant <= 0:
            continue
        if not isinstance(desc, str) or not desc.strip():
            continue
        desc = " ".join(desc.split())
        unit = ws.cell(r, c_unit).value if c_unit else None
        lineas.append({
            "fila": r,
            "cantidad": int(cant) if float(cant).is_integer() else cant,
            "descripcion": desc,
            "unitario": unit if isinstance(unit, (int, float)) else None,
            "es_producto": not any(p in normalizar(desc).lower() for p in PALABRAS_NO_PRODUCTO),
        })
    return lineas


# ----------------------------------------------------------------------------
# 2. Biblioteca local de productos
# ----------------------------------------------------------------------------
class Biblioteca:
    def __init__(self, carpeta="biblioteca"):
        self.dir = Path(carpeta)
        self.dir.mkdir(parents=True, exist_ok=True)
        self._idx = self.dir / "indice.json"
        self.indice: dict = json.loads(self._idx.read_text("utf-8")) if self._idx.exists() else {}

    def _guardar_indice(self):
        self._idx.write_text(json.dumps(self.indice, ensure_ascii=False, indent=2), "utf-8")

    @staticmethod
    def _slug(clave: str) -> str:
        return re.sub(r"[^A-Za-z0-9._-]", "_", clave.strip().upper())

    def carpeta(self, clave: str) -> Path:
        return self.dir / self._slug(clave)

    def archivos(self, clave: str):
        """(foto, ficha): ruta o None si falta."""
        if not clave:
            return None, None
        c = self.carpeta(clave)
        foto = next((p for e in EXT_FOTO if (p := c / f"foto{e}").exists()), None)
        ficha = c / "ficha.pdf"
        return foto, (ficha if ficha.exists() else None)

    def guardar_archivo(self, clave: str, tipo: str, nombre_origen: str, datos: bytes):
        c = self.carpeta(clave)
        c.mkdir(parents=True, exist_ok=True)
        if tipo == "foto":
            ext = Path(nombre_origen).suffix.lower()
            if ext not in EXT_FOTO:
                raise ValueError(f"Formato de imagen no soportado: {ext}")
            for e in EXT_FOTO:  # reemplaza la foto anterior
                (c / f"foto{e}").unlink(missing_ok=True)
            (c / f"foto{ext}").write_bytes(datos)
        elif tipo == "ficha":
            (c / "ficha.pdf").write_bytes(datos)
        else:
            raise ValueError("tipo debe ser 'foto' o 'ficha'")

    def registrar(self, clave: str, nombre: str | None = None, alias: str | None = None,
                  precio_neto: float | None = None):
        """Anota el producto en el índice. El alias permite reconocer la misma
        descripción en futuras cotizaciones aunque no traiga el código; el precio
        neto se recuerda para autocompletar la próxima orden."""
        if not clave:
            return
        clave = clave.strip().upper()
        ficha = self.indice.setdefault(clave, {"nombre": nombre or clave, "aliases": []})
        if precio_neto:
            ficha["precio_neto"] = float(precio_neto)
        if alias and len(normalizar(alias)) >= 6:
            n = normalizar(alias)
            if n not in ficha["aliases"]:
                ficha["aliases"].append(n)
        self._guardar_indice()

    # ---------------------------------------------------------- catálogo editable
    CAMPOS_EDITABLES = ("nombre", "costo_neto", "margen_pct", "notas", "precio_neto")

    def actualizar(self, clave: str, **campos):
        """Cambia nombre, costo neto, margen %, notas o último precio. None borra el dato."""
        ficha = self.indice.get(clave.strip().upper())
        if ficha is None:
            return
        for k, v in campos.items():
            if k not in self.CAMPOS_EDITABLES:
                continue
            if v is None or v == "" or (isinstance(v, float) and math.isnan(v)):
                if k != "nombre":
                    ficha.pop(k, None)
            else:
                ficha[k] = float(v) if k in ("costo_neto", "margen_pct", "precio_neto") else str(v).strip()
        self._guardar_indice()

    def eliminar(self, clave: str):
        """Borra el producto del catálogo junto con su foto y su ficha."""
        clave = clave.strip().upper()
        self.indice.pop(clave, None)
        shutil.rmtree(self.carpeta(clave), ignore_errors=True)
        self._guardar_indice()

    def renombrar(self, viejo: str, nuevo: str, combinar: bool = False) -> str:
        """Cambia el código de un producto (y mueve su carpeta). Si el nuevo ya existe hay que
        pedir combinar=True: se unen los alias y se conserva lo que el nuevo ya tenía."""
        viejo, nuevo = viejo.strip().upper(), nuevo.strip().upper()
        if not nuevo or viejo == nuevo or viejo not in self.indice:
            raise ValueError("Código no válido.")
        if nuevo in self.indice and not combinar:
            raise ValueError("Ya existe un producto con ese código.")
        origen = self.indice.pop(viejo)
        destino = self.indice.setdefault(nuevo, {"nombre": origen.get("nombre", nuevo), "aliases": []})
        destino["aliases"] = sorted(set(destino.get("aliases", [])) | set(origen.get("aliases", []))
                                    | {normalizar(viejo)})
        for k, v in origen.items():
            if k not in ("aliases", "nombre") and k not in destino:
                destino[k] = v
        carp_v, carp_n = self.carpeta(viejo), self.carpeta(nuevo)
        if carp_v.exists():
            carp_n.mkdir(parents=True, exist_ok=True)
            for archivo in carp_v.iterdir():
                if not (carp_n / archivo.name).exists() and not (
                        archivo.stem == "foto" and any((carp_n / f"foto{e}").exists() for e in EXT_FOTO)):
                    shutil.move(str(archivo), carp_n / archivo.name)
            shutil.rmtree(carp_v, ignore_errors=True)
        self._guardar_indice()
        return nuevo

    def precio_sugerido(self, clave: str) -> float | None:
        """Precio neto sugerido = costo × (1 + margen %). None si falta el costo o el margen."""
        ficha = self.indice.get(clave.strip().upper(), {})
        if ficha.get("costo_neto") and ficha.get("margen_pct") is not None:
            return ficha["costo_neto"] * (1 + ficha["margen_pct"] / 100)
        return None

    def filas_catalogo(self) -> list[dict]:
        filas = []
        for clave, f in sorted(self.indice.items()):
            foto, ficha = self.archivos(clave)
            filas.append({"clave": clave, "nombre": f.get("nombre", clave), "foto": bool(foto), "ficha": bool(ficha),
                          "precio_neto": f.get("precio_neto"), "costo_neto": f.get("costo_neto"),
                          "margen_pct": f.get("margen_pct"), "sugerido_neto": self.precio_sugerido(clave),
                          "notas": f.get("notas", "")})
        return filas

    def detectar(self, descripcion: str, solo_conocidos: bool = False) -> str:
        """Clave del producto según la descripción ('' si no se reconoce).

        1) Busca claves/alias ya conocidos en la biblioteca (gana el más largo).
        2) Si no, toma el primer token que parezca código de fabricante
           (mezcla de letras y números: CF283A, TK-3182, T544...).
        """
        dn = normalizar(descripcion)
        mejor, largo = "", 0
        for clave, datos in self.indice.items():
            for token in [clave] + datos.get("aliases", []):
                n = normalizar(token)
                if len(n) >= 3 and n in dn and len(n) > largo:
                    mejor, largo = clave, len(n)
        if mejor or solo_conocidos:
            return mejor
        for tok in re.findall(r"\b[A-Z0-9]+(?:-[A-Z0-9]+)*\b", descripcion.upper()):
            if len(tok) >= 4 and re.search(r"[A-Z]", tok) and re.search(r"\d", tok) \
                    and not _MEDIDA.match(tok):
                return tok
        return ""



def clave_descriptiva(descripcion: str) -> str:
    """Clave legible para productos sin código de fabricante:
    SACA-CORCHETES-PALANCA-METALICO.
    """
    t = unicodedata.normalize(
        "NFKD",
        limpiar_consulta(descripcion)
    ).encode("ascii", "ignore").decode()

    palabras = [
        re.sub(r"[^A-Za-z0-9]", "", w).upper()
        for w in t.split()
    ]

    palabras = [w for w in palabras if w and not w.isdigit()][:4]

    return "-".join(palabras)[:40] or "PRODUCTO"



def resolver_clave(bib: "Biblioteca", descripcion: str, ia: dict | None = None) -> tuple[str, str]:
    """(clave, origen). Prioridad: lo que ya aprendió la biblioteca > la IA (certeza alta o media)
    > código encontrado en el texto > clave descriptiva. Así todo producto queda buscable."""
    if conocido := bib.detectar(descripcion, solo_conocidos=True):
        return conocido, "biblioteca"
    if ia and ia.get("clave") and ia.get("certeza") in ("alta", "media"):
        return ia["clave"].strip().upper(), "ia"
    if texto := bib.detectar(descripcion):
        return texto, "texto"
    return clave_descriptiva(descripcion), "descripcion"


# ----------------------------------------------------------------------------
# 3. Generación del PDF
# ----------------------------------------------------------------------------
def _comando_soffice() -> list[str]:
    cmd = os.environ.get("SOFFICE_CMD")
    if cmd:
        return shlex.split(cmd)
    for nombre in ("soffice", "libreoffice"):
        if ruta := shutil.which(nombre):
            return [ruta]
    for ruta in (r"C:\Program Files\LibreOffice\program\soffice.exe",
                 r"C:\Program Files (x86)\LibreOffice\program\soffice.exe",
                 "/Applications/LibreOffice.app/Contents/MacOS/soffice"):
        if os.path.exists(ruta):
            return [ruta]
    raise RuntimeError("No encontré LibreOffice (gratis en libreoffice.org). "
                       "Instálalo o define la variable SOFFICE_CMD con la ruta de soffice.")


def excel_a_pdf(xlsx, destino, motor="nativo") -> Path:
    """Convierte la cotización a PDF.

    motor="nativo": la dibuja directamente (rápido, sin programas externos ni ventanas).
    motor="libreoffice": usa LibreOffice para un formato idéntico al Excel.
    """
    if motor == "nativo":
        return pdf_cotizacion.render_pdf(xlsx, destino)
    with tempfile.TemporaryDirectory() as tmp:
        copia = Path(tmp) / "cotizacion.xlsx"
        wb = load_workbook(xlsx)
        ws = wb.active
        ws.sheet_properties.pageSetUpPr = PageSetupProperties(fitToPage=True)
        ws.page_setup.fitToWidth = 1
        ws.page_setup.fitToHeight = 0
        wb.save(copia)
        extra = {"creationflags": 0x08000000} if os.name == "nt" else {}  # sin ventana de consola
        perfil = (Path(tmp) / "perfil_lo").as_uri()  # perfil propio: no choca con un LibreOffice abierto
        subprocess.run([*_comando_soffice(), "--headless", "--invisible", "--norestore", "--nolockcheck",
                        "--nologo", "--nofirststartwizard", f"-env:UserInstallation={perfil}",
                        "--convert-to", "pdf", "--outdir", tmp, str(copia)],
                       check=True, capture_output=True, timeout=180, cwd=tmp, **extra)
        shutil.copy(Path(tmp) / "cotizacion.pdf", destino)
    return Path(destino)


def _imagen(ruta, max_w, max_h) -> Image:
    with PILImage.open(ruta) as im:
        if im.mode in ("RGBA", "LA", "P"):  # fondos transparentes -> blanco
            im = im.convert("RGBA")
            fondo = PILImage.new("RGB", im.size, "white")
            fondo.paste(im, mask=im.split()[-1])
            im = fondo
        else:
            im = im.convert("RGB")
        buf = io.BytesIO()
        im.save(buf, "JPEG", quality=88)
        w, h = im.size
    buf.seek(0)
    esc = min(max_w / w, max_h / h)
    return Image(buf, width=w * esc, height=h * esc)


def anexo_fotos(items: list[dict], destino) -> Path:
    """PDF con una grilla de 2 columnas: foto + clave + descripción."""
    estilos = getSampleStyleSheet()
    pie = ParagraphStyle("pie", parent=estilos["Normal"], fontSize=9, leading=11, alignment=1)
    celdas = []
    for it in items:
        foto = it["foto"]
        if foto:
            try:
                img = _imagen(foto, 7.5 * cm, 5.6 * cm)
            except Exception:
                foto = None
        if not foto:
            img = Paragraph("<font color='#999999'>Sin imagen</font>", pie)
        titulo = f"<b>{escape(it['clave'])}</b><br/>" if it["clave"] else ""
        desc = escape(it["descripcion"][:130])
        celdas.append([img, Spacer(1, 6),
                       Paragraph(f"{titulo}{desc}<br/><font color='#666666'>Cantidad cotizada: "
                                 f"{it['cantidad']}</font>", pie)])
    filas = [celdas[i:i + 2] for i in range(0, len(celdas), 2)]
    if filas and len(filas[-1]) == 1:
        filas[-1].append("")
    tabla = Table(filas, colWidths=[9 * cm, 9 * cm], rowHeights=[8.0 * cm] * len(filas))
    tabla.setStyle(TableStyle([
        ("VALIGN", (0, 0), (-1, -1), "TOP"), ("ALIGN", (0, 0), (-1, -1), "CENTER"),
        ("BOX", (0, 0), (-1, -1), 0.5, colors.lightgrey),
        ("INNERGRID", (0, 0), (-1, -1), 0.5, colors.lightgrey),
        ("TOPPADDING", (0, 0), (-1, -1), 10),
    ]))

    def numerar(canvas, doc):
        canvas.setFont("Helvetica", 8)
        canvas.setFillColor(colors.grey)
        canvas.drawRightString(letter[0] - 1.5 * cm, 1 * cm, f"Anexo 1 · Fotografías · pág. {doc.page}")

    doc = SimpleDocTemplate(str(destino), pagesize=letter, leftMargin=1.5 * cm, rightMargin=1.5 * cm,
                            topMargin=1.5 * cm, bottomMargin=1.8 * cm)
    doc.build([Paragraph("Anexo 1 · Fotografías de los productos", estilos["Heading2"]),
               Spacer(1, 8), tabla], onFirstPage=numerar, onLaterPages=numerar)
    return Path(destino)


def generar_pdf_final(xlsx, lineas: list[dict], bib: Biblioteca, salida, motor="nativo") -> dict:
    """Arma: cotización + anexo de fotos + fichas técnicas (una por producto).

    `lineas` son las líneas a incluir, cada una con 'clave' (puede ir vacía).
    Devuelve un informe con lo que faltó (sin clave / sin foto / sin ficha).
    """
    informe = {"sin_clave": [], "sin_foto": [], "sin_ficha": [], "ficha_dudosa": [], "ficha_ilegible": []}
    items, vistos, fichas = [], set(), []
    for ln in lineas:
        clave = ln.get("clave", "").strip().upper()
        if clave and clave in vistos:  # mismo producto en dos líneas: una sola foto/ficha
            continue
        foto, ficha = bib.archivos(clave)
        if not clave:
            informe["sin_clave"].append(ln["descripcion"][:60])
        else:
            vistos.add(clave)
            if not foto:
                informe["sin_foto"].append(clave)
            if not ficha:
                informe["sin_ficha"].append(clave)
            else:
                fichas.append((clave, ficha))
                try:  # ¿la ficha guardada realmente nombra este producto?
                    if es_codigo(clave) and len(menciones_clave(ficha.read_bytes(), [clave])) == 0:
                        informe["ficha_dudosa"].append(clave)
                except Exception:
                    informe["ficha_dudosa"].append(clave)
        items.append({**ln, "clave": clave, "foto": foto})

    with tempfile.TemporaryDirectory() as tmp:
        cot = excel_a_pdf(xlsx, Path(tmp) / "cot.pdf", motor)
        escritor = PdfWriter()
        escritor.append(str(cot), outline_item="Cotización")
        if items:
            escritor.append(str(anexo_fotos(items, Path(tmp) / "fotos.pdf")),
                            outline_item="Anexo 1 · Fotografías")
        for clave, ficha in fichas:
            try:
                escritor.append(str(ficha), outline_item=f"Ficha técnica · {clave}")
            except Exception:
                informe["ficha_ilegible"].append(clave)
        with open(salida, "wb") as f:
            escritor.write(f)
    informe["paginas"] = len(PdfReader(str(salida)).pages)
    return informe


# ----------------------------------------------------------------------------
# 4. Orden en bruto -> cotización con formato (usa tu plantilla de Excel)
# ----------------------------------------------------------------------------
PLANTILLA = Path(__file__).with_name("plantilla_cotizacion.xlsx")
IVA = 0.19
_QTY_INICIO = re.compile(r"^(\d+(?:[.,]\d+)?)\s*(?:x|×|un\.?|unid\.?|unidades?|u\.?)?\s*[:\-–.)]*\s+(\S.*)$", re.I)
_QTY_FINAL = re.compile(r"^(\S.*?)[\s,;:\-–]*[(\[]?\s*(?:x|×|cant(?:idad)?[:.]?)\s*(\d+)\s*"
                        r"(?:un\w*\.?|u\.?)?[)\]]?$", re.I)
_QTY_PAREN = re.compile(r"^(\S.*?)[\s,;:\-–]*\(\s*(\d+)\s*(?:un\w*\.?|u\.?)\s*\)$", re.I)
_SINONIMOS_CANT = {"CANTIDAD", "CANT", "QTY", "UNIDADES", "UNID", "CANTIDADES"}
_SINONIMOS_DESC = {"DESCRIPCION", "PRODUCTO", "PRODUCTOS", "DETALLE", "ITEM", "ARTICULO", "NOMBRE", "GLOSA"}


def limpiar_descripcion(texto: str) -> str:
    """Espacios normalizados y primera letra en mayúscula."""
    t = " ".join(str(texto).split())
    return t[:1].upper() + t[1:]


def _num(valor):
    try:
        return float(str(valor).replace(".", "").replace(",", ".")) if isinstance(valor, str) else float(valor)
    except (TypeError, ValueError):
        return None


def _cantidad(n: float):
    return int(n) if float(n).is_integer() else n


def parsear_texto_orden(texto: str) -> list[dict]:
    """Convierte texto pegado (una línea por producto) en [{cantidad, descripcion}].

    Acepta: '10 Cartucho CF283A', '10x CF283A', 'Cartucho CF283A x10',
    'Cartucho CF283A (10 un)' y filas pegadas desde Excel (separadas por tab).
    Si no encuentra cantidad usa 1 (se corrige en la tabla de revisión).
    """
    salida = []
    for linea in texto.splitlines():
        s = linea.strip().lstrip("•-*·").strip()
        if not s:
            continue
        if "\t" in s:
            partes = [p.strip() for p in s.split("\t") if p.strip()]
            nums = [(i, _num(p)) for i, p in enumerate(partes) if _num(p) is not None]
            textos = [p for p in partes if _num(p) is None]
            if nums and textos:
                salida.append({"cantidad": _cantidad(nums[0][1]), "descripcion": limpiar_descripcion(max(textos, key=len))})
                continue
        if m := _QTY_INICIO.match(s):
            salida.append({"cantidad": _cantidad(_num(m.group(1)) or 1), "descripcion": limpiar_descripcion(m.group(2))})
        elif (m := _QTY_FINAL.match(s)) or (m := _QTY_PAREN.match(s)):
            salida.append({"cantidad": int(m.group(2)), "descripcion": limpiar_descripcion(m.group(1))})
        else:
            salida.append({"cantidad": 1, "descripcion": limpiar_descripcion(s)})
    return salida


def leer_orden_archivo(ruta) -> list[dict]:
    """Lee una orden en bruto desde Excel o CSV (con o sin encabezados)."""
    ruta = Path(ruta)
    if ruta.suffix.lower() == ".csv":
        crudo = ruta.read_bytes().decode("utf-8-sig", errors="replace")
        dialecto = csv.Sniffer().sniff(crudo[:2000], delimiters=",;\t") if crudo.strip() else csv.excel
        filas = [list(f) for f in csv.reader(io.StringIO(crudo), dialecto)]
    else:
        filas = [[c.value for c in fila] for fila in load_workbook(ruta, data_only=True).active.iter_rows()]

    col_cant = col_desc = None
    for fila in filas[:15]:  # ¿hay encabezados?
        etiquetas = {normalizar(v): i for i, v in enumerate(fila) if isinstance(v, str)}
        c = next((i for k, i in etiquetas.items() if k in _SINONIMOS_CANT), None)
        d = next((i for k, i in etiquetas.items() if k in _SINONIMOS_DESC), None)
        if c is not None and d is not None:
            col_cant, col_desc = c, d
            break

    salida = []
    for fila in filas:
        textos = [str(v).strip() for v in fila if isinstance(v, str) and v.strip() and _num(v) is None]
        if col_desc is not None:
            desc = fila[col_desc] if col_desc < len(fila) else None
            cant = _num(fila[col_cant]) if col_cant < len(fila) else None
            if isinstance(desc, str) and desc.strip() and cant and cant > 0:
                salida.append({"cantidad": _cantidad(cant), "descripcion": limpiar_descripcion(desc)})
            continue
        if not textos:
            continue
        nums = [_num(v) for v in fila if not isinstance(v, str) and _num(v) is not None and _num(v) > 0]
        desc = max(textos, key=len)
        if nums:
            salida.append({"cantidad": _cantidad(nums[0]), "descripcion": limpiar_descripcion(desc)})
        else:
            salida += parsear_texto_orden(desc)
    return salida


def neto_desde(precio: float, con_iva: bool) -> float:
    """Precio neto unitario a partir de lo que el usuario ingresó."""
    return precio / (1 + IVA) if con_iva else float(precio)


def precio_para_mostrar(precio_neto: float, con_iva: bool) -> int:
    return round(precio_neto * (1 + IVA)) if con_iva else round(precio_neto)


def crear_cotizacion_xlsx(lineas: list[dict], cliente: dict, salida, plantilla=None) -> Path:
    """Llena la plantilla con las líneas y los datos del cliente.

    lineas : [{cantidad, descripcion, precio_neto (o None)}]
    cliente: nombre, rut, direccion, ciudad, telefono, fecha (date)
    La plantilla trae 19 filas; si hay más productos se agregan filas
    desplazando los totales y la forma de pago hacia abajo.
    Totales: Neto = suma de líneas, IVA = 19 %, Total = Neto + IVA.
    """
    wb = load_workbook(plantilla or PLANTILLA)
    ws = wb.active
    PRIMERA, ULTIMA = 20, 38
    extra = max(0, len(lineas) - (ULTIMA - PRIMERA + 1))
    if extra:  # baja el pie (última fila de la tabla, totales y forma de pago)
        ws.unmerge_cells(f"C{ULTIMA}:F{ULTIMA}")
        ws.move_range(f"A{ULTIMA}:L{ws.max_row}", rows=extra)
        for r in range(ULTIMA, ULTIMA + extra):
            for col in range(2, 9):
                ws.cell(r, col)._style = copy(ws.cell(ULTIMA - 1, col)._style)
        for r in range(ULTIMA, ULTIMA + extra + 1):
            ws.merge_cells(start_row=r, start_column=3, end_row=r, end_column=6)
    ultima = ULTIMA + extra

    # datos del cliente (etiquetas B11:B14 y F11:F12 de la plantilla)
    ws["C11"], ws["C12"], ws["C13"], ws["C14"] = (cliente.get(k, "") or None for k in
                                                  ("nombre", "direccion", "ciudad", "telefono"))
    ws["G11"] = cliente.get("rut") or None
    ws["G12"] = cliente.get("fecha") or date.today()
    if cliente.get("numero"):  # N° correlativo, bajo el RUT y la fecha
        ws["F13"].value = "N° Cotización:"
        ws["F13"]._style = copy(ws["F12"]._style)
        ws["G13"].value = cliente["numero"]
        ws["G13"]._style = copy(ws["G12"]._style)
        ws["G13"].number_format = "@"

    ancho_chars = sum((ws.column_dimensions[c].width or 8.43) for c in "CDEF")
    for i, r in enumerate(range(PRIMERA, ultima + 1)):
        ln = lineas[i] if i < len(lineas) else None
        ws.cell(r, 2).value = ln["cantidad"] if ln else None
        ws.cell(r, 3).value = ln["descripcion"] if ln else None
        ws.cell(r, 3).alignment = Alignment(wrap_text=True, vertical="center")
        for col in (2, 3, 7, 8):  # misma letra en todas las filas (la plantilla mezclaba negritas)
            ws.cell(r, col).font = copy(ws.cell(PRIMERA + 1, col).font)
        ws.cell(r, 7).value = ln.get("precio_neto") if ln else None
        ws.cell(r, 8).value = f"=G{r}*B{r}"
        if ln:  # las celdas combinadas no se autoajustan: calcula el alto
            lineas_texto = max(1, math.ceil(len(ln["descripcion"]) / (ancho_chars * 0.95)))
            ws.row_dimensions[r].height = 15 * lineas_texto + 2
        else:
            ws.row_dimensions[r].height = None

    t_neto, t_iva, t_total = ultima + 1, ultima + 2, ultima + 3
    ws[f"H{t_neto}"] = f"=ROUND(SUM(H{PRIMERA}:H{ultima}),0)"
    ws[f"H{t_iva}"] = f"=ROUND(SUM(H{PRIMERA}:H{ultima})*{IVA},0)"
    ws[f"H{t_total}"] = f"=H{t_neto}+H{t_iva}"
    for celda in ("K22", "K24", "K26"):  # auxiliares de la plantilla original
        ws[celda] = None

    ws.sheet_properties.pageSetUpPr = PageSetupProperties(fitToPage=True)
    ws.page_setup.fitToWidth, ws.page_setup.fitToHeight = 1, 0
    wb.save(salida)
    return Path(salida)
