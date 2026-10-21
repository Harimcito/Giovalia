"""PDF de la cotización generado directamente (sin LibreOffice ni ventanas externas).

Lee la cotización en Excel (encabezado de la empresa, logo, datos del cliente, líneas,
totales y forma de pago) y la dibuja con el mismo orden y estilo de la plantilla.
"""
import io
import re
import unicodedata
from datetime import date, datetime
from pathlib import Path
from xml.sax.saxutils import escape

from openpyxl import load_workbook
from reportlab.lib import colors
from reportlab.lib.pagesizes import letter
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import cm
from reportlab.platypus import (Image, KeepTogether, Paragraph, SimpleDocTemplate, Spacer, Table,
                                TableStyle)

IVA = 0.19
AZUL = colors.HexColor("#0000FF")


def _norm(texto) -> str:
    t = unicodedata.normalize("NFKD", str(texto)).encode("ascii", "ignore").decode()
    return re.sub(r"[^A-Z0-9]", "", t.upper())


def _num(v):
    return float(v) if isinstance(v, (int, float)) and not isinstance(v, bool) else None


def clp(valor) -> str:
    """Pesos chilenos con punto de miles: $ 1.234.567"""
    return "$ " + f"{round(valor):,}".replace(",", ".")


def _texto_celda(v) -> str:
    if v is None:
        return ""
    if isinstance(v, (datetime, date)):
        return v.strftime("%d/%m/%Y")
    return str(v).strip()


def leer_para_pdf(xlsx) -> dict:
    """Extrae del Excel todo lo que se dibuja en el PDF."""
    wb = load_workbook(xlsx, data_only=True)  # usa los valores calculados si el archivo los trae
    ws = wb.active

    fila_enc, cols = None, {}
    for fila in ws.iter_rows():
        etiquetas = {_norm(c.value): c.column for c in fila if isinstance(c.value, str)}
        if "CANTIDAD" in etiquetas and "DESCRIPCION" in etiquetas:
            fila_enc, cols = fila[0].row, etiquetas
            break
    if fila_enc is None:
        raise ValueError("No encontré los encabezados 'Cantidad' y 'Descripción' en la hoja.")
    c_cant, c_desc = cols["CANTIDAD"], cols["DESCRIPCION"]
    c_unit, c_total = cols.get("PUNITARIO"), cols.get("TOTAL")

    # zona del cliente: desde la primera etiqueta "Algo:" hasta la tabla
    fila_cliente = next((c.row for fila in ws.iter_rows(max_row=fila_enc - 1) for c in fila
                         if isinstance(c.value, str) and c.value.strip().endswith(":")), fila_enc)
    empresa = [str(c.value).strip() for fila in ws.iter_rows(max_row=fila_cliente - 1)
               for c in fila if isinstance(c.value, str) and c.value.strip()]

    pares = []  # (columna, etiqueta, valor)
    for fila in ws.iter_rows(min_row=fila_cliente, max_row=fila_enc - 1):
        celdas = [c for c in fila if c.value is not None]
        for i, c in enumerate(celdas):
            if isinstance(c.value, str) and c.value.strip().endswith(":"):
                sig = celdas[i + 1] if i + 1 < len(celdas) else None
                valor = "" if sig is None or (isinstance(sig.value, str) and sig.value.strip().endswith(":")) \
                    else _texto_celda(sig.value)
                pares.append((c.column, c.value.strip(), valor))
    col_izq = min((p[0] for p in pares), default=0)
    izquierda = [(e, v) for col, e, v in pares if col == col_izq]
    derecha = [(e, v) for col, e, v in pares if col != col_izq]

    lineas, ultima = [], fila_enc
    for r in range(fila_enc + 1, ws.max_row + 1):
        cant, desc = _num(ws.cell(r, c_cant).value), ws.cell(r, c_desc).value
        if cant and cant > 0 and isinstance(desc, str) and desc.strip():
            unit = _num(ws.cell(r, c_unit).value) if c_unit else None
            total = _num(ws.cell(r, c_total).value) if c_total else None
            if total is None and unit is not None:
                total = unit * cant
            lineas.append({"cantidad": int(cant) if cant.is_integer() else cant,
                           "descripcion": " ".join(desc.split()), "unitario": unit, "total": total})
            ultima = r

    # totales: etiquetas "Total Neto", "I.V.A.", "Total" bajo la tabla (valor en la celda de al lado)
    totales, ultima_fila_totales = [], ultima
    for r in range(fila_enc + 1, ws.max_row + 1):
        for c in ws[r]:
            if isinstance(c.value, str) and _norm(c.value) in ("TOTALNETO", "IVA", "TOTAL"):
                totales.append([c.value.strip(), _num(ws.cell(r, c.column + 1).value)])
                ultima_fila_totales = max(ultima_fila_totales, r)
    if any(v is None for _, v in totales):  # archivo sin valores calculados: neto + IVA 19 %
        suma = sum(ln["total"] or 0 for ln in lineas)
        neto = round(suma)
        calculados = {"TOTALNETO": neto, "IVA": round(suma * IVA), "TOTAL": neto + round(suma * IVA)}
        totales = [[e, calculados[_norm(e)]] for e, _ in totales]

    pago = [str(ws.cell(r, c_cant).value).strip() for r in range(ultima_fila_totales + 1, ws.max_row + 1)
            if ws.cell(r, c_cant).value not in (None, "")]
    logo = ws._images[0]._data() if ws._images else None
    return {"empresa": empresa, "izquierda": izquierda, "derecha": derecha, "lineas": lineas,
            "totales": totales, "pago": pago, "logo": logo, "tiene_precios": c_unit is not None}


def render_pdf(xlsx, destino) -> Path:
    d = leer_para_pdf(xlsx)
    est = getSampleStyleSheet()
    base = ParagraphStyle("b", parent=est["Normal"], fontName="Helvetica", fontSize=9, leading=11)
    negrita = ParagraphStyle("n", parent=base, fontName="Helvetica-Bold")
    centro = ParagraphStyle("c", parent=base, alignment=1)
    centro_b = ParagraphStyle("cb", parent=negrita, alignment=1)
    derecha = ParagraphStyle("d", parent=base, alignment=2)
    empresa_st = ParagraphStyle("e", parent=negrita, textColor=AZUL, fontSize=8.5, leading=10.5)

    margen = 1.6 * cm
    ancho = letter[0] - 2 * margen
    story = []

    # encabezado: logo + datos de la empresa
    logo = Spacer(1, 1)
    if d["logo"]:
        from PIL import Image as PILImage
        with PILImage.open(io.BytesIO(d["logo"])) as im:
            w, h = im.size
        logo = Image(io.BytesIO(d["logo"]), width=6.3 * cm, height=6.3 * cm * h / w)
    texto_empresa = [Paragraph(escape(t), empresa_st) for t in d["empresa"]]
    cab = Table([[logo, texto_empresa]], colWidths=[7.2 * cm, ancho - 7.2 * cm])
    cab.setStyle(TableStyle([("VALIGN", (0, 0), (-1, -1), "MIDDLE"), ("LINEBELOW", (0, 0), (-1, -1), 1, colors.black),
                             ("BOTTOMPADDING", (0, 0), (-1, -1), 8)]))
    story += [cab, Spacer(1, 10)]

    # datos del cliente
    filas = []
    for i in range(max(len(d["izquierda"]), len(d["derecha"]))):
        e1, v1 = d["izquierda"][i] if i < len(d["izquierda"]) else ("", "")
        e2, v2 = d["derecha"][i] if i < len(d["derecha"]) else ("", "")
        filas.append([Paragraph(escape(e1), negrita), Paragraph(escape(v1), base),
                      Paragraph(escape(e2), negrita), Paragraph(escape(v2), base)])
    if filas:
        cli = Table(filas, colWidths=[2.3 * cm, ancho * 0.5 - 2.3 * cm, 3.3 * cm, ancho * 0.5 - 3.3 * cm])
        cli.setStyle(TableStyle([("BOX", (0, 0), (-1, -1), 1.5, colors.black), ("VALIGN", (0, 0), (-1, -1), "TOP"),
                                 ("TOPPADDING", (0, 0), (-1, -1), 1.5), ("BOTTOMPADDING", (0, 0), (-1, -1), 1.5)]))
        story += [cli, Spacer(1, 10)]

    # tabla de productos
    w_cant, w_unit, w_tot = 2.2 * cm, 3.0 * cm, 3.4 * cm
    w_desc = ancho - w_cant - w_unit - w_tot
    filas = [[Paragraph("Cantidad", centro_b), Paragraph("Descripción", centro_b),
              Paragraph("P. Unitario", centro_b), Paragraph("TOTAL", centro_b)]]
    for ln in d["lineas"]:
        filas.append([Paragraph(str(ln["cantidad"]), centro), Paragraph(escape(ln["descripcion"]), base),
                      Paragraph("" if ln["unitario"] is None else clp(ln["unitario"]), derecha),
                      Paragraph("" if ln["total"] is None else clp(ln["total"]), derecha)])
    tabla = Table(filas, colWidths=[w_cant, w_desc, w_unit, w_tot], repeatRows=1)
    tabla.setStyle(TableStyle([
        ("BOX", (0, 0), (-1, -1), 1, colors.black), ("LINEBELOW", (0, 0), (-1, 0), 1, colors.black),
        ("LINEAFTER", (0, 0), (2, -1), 0.6, colors.black), ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
        ("TOPPADDING", (0, 0), (-1, 0), 6), ("BOTTOMPADDING", (0, 0), (-1, 0), 6),
        ("TOPPADDING", (0, 1), (-1, -1), 4), ("BOTTOMPADDING", (0, 1), (-1, -1), 4),
    ]))
    story.append(tabla)

    # totales (alineados bajo las dos últimas columnas)
    if d["totales"]:
        tot = Table([[Paragraph(escape(e), negrita), Paragraph(clp(v), ParagraphStyle("t", parent=negrita, alignment=2))]
                     for e, v in d["totales"]], colWidths=[w_unit, w_tot], hAlign="RIGHT")
        tot.setStyle(TableStyle([("BOX", (0, 0), (-1, -1), 1, colors.black),
                                 ("LINEBELOW", (0, 0), (-1, -2), 0.5, colors.black),
                                 ("TOPPADDING", (0, 0), (-1, -1), 3), ("BOTTOMPADDING", (0, 0), (-1, -1), 3)]))
        story.append(KeepTogether([Spacer(1, 4), tot]))

    # forma de pago
    if d["pago"]:
        caja = Table([[Paragraph(escape(t), base)] for t in d["pago"]], colWidths=[7.5 * cm], hAlign="LEFT")
        caja.setStyle(TableStyle([("BOX", (0, 0), (-1, -1), 1.5, colors.black),
                                  ("TOPPADDING", (0, 0), (-1, -1), 1), ("BOTTOMPADDING", (0, 0), (-1, -1), 1)]))
        story.append(KeepTogether([Spacer(1, 14), caja]))

    SimpleDocTemplate(str(destino), pagesize=letter, leftMargin=margen, rightMargin=margen,
                      topMargin=1.4 * cm, bottomMargin=1.4 * cm).build(story)
    return Path(destino)
