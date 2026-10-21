"""Cotizaciones con fotos y fichas técnicas. Ejecutar con:  streamlit run app.py"""
import re
import tempfile
import time
from datetime import date, datetime
from pathlib import Path
from copy import copy

from openpyxl import load_workbook
from urllib.parse import quote_plus

import pandas as pd
import streamlit as st

import buscador as bs
import cotizador as ct
import historial as hs
import lector_ia as li
import pdf_cotizacion as pc
import revisor as rv

APP_DIR = Path(__file__).parent
E = st.session_state

# ------------------------------------------------------------- DISEÑO "tinta en agua"
ANIMAR_FONDO = True   # pon False si prefieres el fondo quieto
TEMA_BASE = {"base": "light", "primaryColor": "#2A7090", "backgroundColor": "#F5F9FB",
             "secondaryBackgroundColor": "#E3EFF5", "textColor": "#17394F", "font": "sans serif"}
TEMA_EXTRA = {"baseRadius": "large", "buttonRadius": "large", "borderColor": "#C9DEE8",
              "linkColor": "#2A7090", "dataframeHeaderBackgroundColor": "#DCEBF3"}


def preparar_tema() -> bool:
    """Instala el diseño claro la primera vez. Devuelve True si esta ventana ya lo está usando."""
    ya_activo = st.get_option("theme.base") == "light"
    try:  # queda guardado para las próximas veces que abras el programa
        cfg = APP_DIR / ".streamlit" / "config.toml"
        texto = cfg.read_text("utf-8") if cfg.exists() else ""
        if "[theme]" not in texto:
            cfg.parent.mkdir(exist_ok=True)
            bloque = "[theme]\n" + "\n".join(f'{k} = "{v}"' for k, v in TEMA_BASE.items()) + "\n"
            cfg.write_text(texto + ("\n" if texto else "") + bloque, "utf-8")
    except OSError:
        pass
    for k, v in {**TEMA_BASE, **TEMA_EXTRA}.items():  # vale desde la próxima actualización de la página
        try:
            st._config.set_option(f"theme.{k}", v)
        except Exception:
            pass
    return ya_activo


TEMA_ACTIVO = preparar_tema()

st.set_page_config(page_title="Cotizaciones con fotos y fichas", page_icon="💧",
                   layout="wide", initial_sidebar_state="collapsed")

CSS_BASE = """
html, body, .stApp { font-size: 18px; }
.block-container { max-width: 1280px; padding-top: 2.4rem; }
h2 { font-size: 1.6rem !important; }
.stButton > button, .stDownloadButton > button, .stLinkButton > a {
    font-size: 1.05rem; padding: 0.65rem 1.2rem; border-radius: 14px; }
.stAppDeployButton, [data-testid="stAppDeployButton"] { display: none; }
.pasos { display: flex; gap: 8px; flex-wrap: wrap; margin: 4px 0 18px 0; }
.paso { padding: 8px 16px; border-radius: 999px; border: 1px solid rgba(128,128,128,.45);
        opacity: .7; font-size: .95rem; }
.paso b { margin-right: 6px; }
.paso.activo { opacity: 1; border: 2px solid #ff4b4b; background: rgba(255,75,75,.18); font-weight: 600; }
.paso.hecho { opacity: 1; border-color: rgba(0,190,100,.7); background: rgba(0,190,100,.15); }
"""

CSS_TINTA = """
.stApp {
  background:
    radial-gradient(60vw 50vh at 5% 0%, rgba(86,170,205,.34), transparent 65%),
    radial-gradient(50vw 45vh at 100% 8%, rgba(158,138,222,.26), transparent 65%),
    radial-gradient(55vw 50vh at 85% 100%, rgba(110,205,190,.30), transparent 65%),
    radial-gradient(45vw 40vh at 0% 95%, rgba(245,165,190,.22), transparent 65%),
    linear-gradient(180deg, #F7FBFD 0%, #EEF6FA 100%);
  background-size: 130% 130%, 130% 130%, 130% 130%, 130% 130%, 100% 100%;
  background-attachment: fixed;
}
[data-testid="stHeader"] { background: transparent; }
[data-testid="stSidebar"] > div:first-child { background: rgba(255,255,255,.6); backdrop-filter: blur(10px); }
h1, h2, h3 { color: #17394F; letter-spacing: -0.01em; }
h1.titulo { font-size: 2.1rem; font-weight: 700; margin: 0 0 .4rem 0; padding: 0;
  background: linear-gradient(90deg, #17394F 0%, #2A7090 55%, #7A6BC4 100%);
  -webkit-background-clip: text; background-clip: text; color: transparent; }
.gota { display: inline-block; width: .8em; height: .8em; margin-right: .45rem; vertical-align: -0.05em;
  border-radius: 62% 38% 55% 45% / 55% 45% 60% 40%;
  background: radial-gradient(circle at 30% 30%, #A9E1F0, #2A7090 58%, #6C5CB5); }
button[kind="primary"], [data-testid="stBaseButton-primary"] {
  background: linear-gradient(135deg, #2A7090 0%, #4F8FB5 55%, #8A7FCB 100%); border: none; color: #fff;
  box-shadow: 0 6px 18px rgba(42,112,144,.28); }
button[kind="primary"]:hover, [data-testid="stBaseButton-primary"]:hover { filter: brightness(1.08); color: #fff; }
button[kind="primary"]:disabled, [data-testid="stBaseButton-primary"]:disabled { opacity: .45; box-shadow: none; }
button[kind="secondary"], [data-testid="stBaseButton-secondary"] {
  background: rgba(255,255,255,.78); border: 1px solid rgba(42,112,144,.35); color: #17394F; }
button[kind="secondary"]:hover, [data-testid="stBaseButton-secondary"]:hover { background: #fff; border-color: #2A7090; }
[data-testid="stVerticalBlockBorderWrapper"]:has(.tarjeta-marca):not(:has([data-testid="stVerticalBlockBorderWrapper"] .tarjeta-marca)) {
  background: rgba(255,255,255,.72); backdrop-filter: blur(8px); border: 1px solid rgba(255,255,255,.95);
  border-radius: 22px; box-shadow: 0 10px 34px rgba(50,100,135,.15); }
[data-testid="stAlert"] { border-radius: 16px; }
[data-baseweb="input"] > div, [data-baseweb="textarea"], [data-baseweb="select"] > div { border-radius: 14px; }
.paso { background: rgba(255,255,255,.62); border-color: rgba(42,112,144,.25); color: #17394F; }
.paso.activo { background: linear-gradient(135deg, #2A7090, #8A7FCB); color: #fff; border: none;
  box-shadow: 0 6px 16px rgba(42,112,144,.28); }
.paso.hecho { background: rgba(110,205,190,.38); border-color: rgba(60,160,150,.6); }
"""
CSS_ANIMACION = """
@keyframes tinta { 0% { background-position: 0% 0%, 100% 0%, 60% 100%, 0% 100%, 0 0; }
                   100% { background-position: 18% 12%, 82% 18%, 70% 82%, 12% 88%, 0 0; } }
.stApp { animation: tinta 80s ease-in-out infinite alternate; }
@media (prefers-reduced-motion: reduce) { .stApp { animation: none; } }
"""
st.markdown("<style>" + CSS_BASE + (CSS_TINTA + (CSS_ANIMACION if ANIMAR_FONDO else "") if TEMA_ACTIVO else "")
            + "</style>", unsafe_allow_html=True)

E.setdefault("paso", 1)

COL_INCLUIR, COL_CANT, COL_PROD, COL_COD, COL_ORIGEN, COL_PRECIO = (
    "Incluir en anexo", "Cantidad", "Producto", "Código", "¿Cómo lo identifiqué?", "Precio unitario neto")
ORIGEN = {"biblioteca": "📚 Ya lo conocía", "ia": "🧠 Lo deduje con IA", "texto": "🔤 Estaba en el texto",
          "descripcion": "✏️ Lo armé con el nombre", "": ""}

def _es_vacio(v) -> bool:
    try:
        return v is None or bool(pd.isna(v))
    except (TypeError, ValueError):
        return v is None


def _json_seguro(valor):
    """Convierte valores de pandas/Path/fecha a tipos compatibles con JSON."""
    if isinstance(valor, pd.DataFrame):
        return [_json_seguro(x) for x in valor.to_dict("records")]
    if isinstance(valor, dict):
        return {str(k): _json_seguro(v) for k, v in valor.items()}
    if isinstance(valor, (list, tuple)):
        return [_json_seguro(v) for v in valor]
    if isinstance(valor, (datetime, date)):
        return valor.isoformat()
    if isinstance(valor, Path):
        return str(valor)
    if _es_vacio(valor):
        return None
    if hasattr(valor, "item"):
        try:
            return _json_seguro(valor.item())
        except (ValueError, TypeError):
            pass
    if isinstance(valor, (str, int, float, bool)) or valor is None:
        return valor
    return str(valor)


def iniciar_nueva_cotizacion():
    """Inicia una sesión nueva sin borrar los borradores guardados en SQLite."""
    E.clear()
    E["_forzar_nuevo_borrador"] = True


def restaurar_borrador(registro: dict):
    """Restaura el estado de trabajo persistido para continuar tras un reinicio."""
    datos = registro.get("datos") or {}
    estado = datos.get("_app_state") or {}
    E["cot_id"] = registro["id"]
    E["cot_numero"] = registro["numero"]
    E["num_cot_edit"] = registro["numero"]
    if estado:
        for clave, valor in estado.items():
            if clave in ("tabla_edit", "tabla"):
                E[clave] = pd.DataFrame(valor or [], columns=estado.get("tabla_columns"))
            elif clave == "orden_df":
                E[clave] = pd.DataFrame(valor or [], columns=["Cantidad", "Producto", "Precio unitario"])
            elif clave == "ruta":
                ruta = Path(valor) if valor else None
                if ruta and ruta.exists():
                    E["ruta"] = ruta
            elif clave == "c_fecha":
                try:
                    E[clave] = date.fromisoformat(str(valor)[:10])
                except ValueError:
                    E[clave] = date.today()
            elif clave not in ("pdf_bytes", "pdf_nombre", "informe"):
                E[clave] = valor
        E.setdefault("paso", 1)
        E.setdefault("vista", "🧾 Nueva cotización")
    else:
        cliente = datos.get("cliente") or {}
        E.update(c_nombre=cliente.get("nombre", registro.get("cliente", "")),
                 c_rut=cliente.get("rut", registro.get("rut", "")),
                 c_dir=cliente.get("direccion", ""), c_ciudad=cliente.get("ciudad", ""),
                 c_tel=cliente.get("telefono", ""), c_email=cliente.get("email", ""),
                 c_fecha=date.today(), paso=1, vista="🧾 Nueva cotización")
        if registro.get("xlsx") and Path(registro["xlsx"]).exists():
            E["_reanudar_xlsx"] = registro["xlsx"]
    for clave in list(E.keys()):
        if clave.startswith(("cand_foto_", "cand_ficha_", "diag_foto_", "diag_ficha_")):
            E.pop(clave, None)


def actualizar_numero_en_excel(ruta: Path, numero: str):
    """Actualiza la celda contigua a la etiqueta N° Cotización, si existe."""
    if not ruta or not ruta.exists():
        return
    try:
        wb = load_workbook(ruta)
        ws = wb.active
        for fila in ws.iter_rows():
            for celda in fila:
                if isinstance(celda.value, str) and ct.normalizar(celda.value.rstrip(":")) in (
                    "NCOTIZACION", "NUMEROCOTIZACION", "NUMERODECOTIZACION"
                ):
                    destino = ws.cell(celda.row, celda.column + 1)
                    destino.value = numero
                    destino.number_format = "@"
                    wb.save(ruta)
                    wb.close()
                    return
        wb.close()
    except Exception:
        return


# ============================================================== AJUSTES (barra lateral)
with st.sidebar:
    st.header("⚙️ Ajustes")
    carpeta = st.text_input("Carpeta donde guardo las fotos y fichas", "biblioteca",
                            help="No hace falta cambiarla.")
    bib = ct.Biblioteca(carpeta)
    hist = hs.Historial(carpeta)
    if not E.get("cot_id") or hist.obtener(E.get("cot_id")) is None:
        forzar_nuevo = bool(E.pop("_forzar_nuevo_borrador", False))
        borrador = None if forzar_nuevo else hist.borrador_reciente()
        if borrador:
            restaurar_borrador(borrador)
        else:
            cid, numero = hist.crear({}, [], {"neto": 0, "iva": 0, "total": 0},
                                     con_iva=True, fecha=date.today(), origen="pedido")
            E.update(cot_id=cid, cot_numero=numero, num_cot_edit=numero,
                     paso=1, vista="🧾 Nueva cotización")
    else:
        registro_activo = hist.obtener(E["cot_id"])
        E.setdefault("cot_numero", registro_activo["numero"])
        E.setdefault("num_cot_edit", registro_activo["numero"])
    st.caption(f"📦 {len(bib.indice)} productos guardados")
    api_key, origen_key = li.buscar_api_key(APP_DIR)
    if api_key:
        st.caption(f"🔑 Clave de OpenAI lista ({origen_key})")
    else:
        api_key = st.text_input("Clave de OpenAI", type="password",
                                help="Sirve para entender mejor los productos y leer pedidos en PDF o foto.")
        if api_key and st.button("💾 Guardar esta clave"):
            li.guardar_api_key(APP_DIR, api_key)
            st.success("Guardada. La próxima vez se carga sola.")
    with st.expander("Opciones avanzadas"):
        usar_lo = st.checkbox("Usar LibreOffice para el PDF", value=False,
                              help="Dejar desmarcado. Solo para un formato idéntico al Excel; "
                                   "puede abrir un cuadro de impresora.")
    with st.expander("Numeración de cotizaciones"):
        fmt = st.text_input("Formato del número", value=hist.config("formato_numero", hs.FORMATO_DEFECTO),
                            help="{n} = número correlativo, {aa} = año con 2 cifras, {aaaa} = año completo. "
                                 "Ejemplo: {n:04d}-COT{aa}")
        sig = st.number_input("Próximo número", min_value=1, step=1, value=hist.siguiente_correlativo())
        st.caption(f"Se verá así: {hist.formatear(int(sig), date.today(), fmt)}")
        if st.button("Guardar numeración"):
            hist.set_config("formato_numero", fmt)
            hist.set_config("siguiente_correlativo", int(sig))
            st.success("Guardado.")
    with st.expander("🔢 Número de esta cotización", expanded=False):
        st.text_input("Número editable", key="num_cot_edit",
                      disabled=("pdf_bytes" in E), help="Puedes cambiarlo antes de generar el PDF.")
        if st.button("Aplicar número", key="aplicar_num_cot", disabled=("pdf_bytes" in E)):
            try:
                numero_nuevo = hist.cambiar_numero(int(E["cot_id"]), str(E.get("num_cot_edit", "")).strip())
                E["cot_numero"] = numero_nuevo
                registro_num = hist.obtener(int(E["cot_id"]))
                if registro_num and registro_num.get("xlsx"):
                    E["ruta"] = Path(registro_num["xlsx"])
                    actualizar_numero_en_excel(E["ruta"], numero_nuevo)
                    hist.guardar_archivos(int(E["cot_id"]), xlsx=E["ruta"])
                st.success(f"Número guardado: {numero_nuevo}")
            except ValueError as exc:
                st.error(str(exc))
    st.divider()
    st.button("🔄 Empezar una cotización nueva", on_click=iniciar_nueva_cotizacion)


# ============================================================== FUNCIONES COMUNES
VISTAS = ["🧾 Nueva cotización", "📚 Historial", "📦 Mis productos"]
OPCIONES_PRECIO = ["Incluyen IVA (lo que paga el cliente)", "Son netos (sin IVA)"]
E.setdefault("vista", VISTAS[0])


def editor(df, **kw):
    """Tabla editable que se estira para mostrar todas sus filas, sin barra de desplazamiento interna."""
    try:
        return st.data_editor(df, height="content", **kw)
    except Exception:  # versiones antiguas de Streamlit: calcula el alto a mano
        extra = 1 if kw.get("num_rows") == "dynamic" else 0
        return st.data_editor(df, height=min(35 * (len(df) + 1 + extra) + 8, 2400), **kw)


def _cliente_actual() -> dict:
    return {
        "nombre": E.get("c_nombre", ""),
        "rut": E.get("c_rut", ""),
        "direccion": E.get("c_dir", ""),
        "ciudad": E.get("c_ciudad", ""),
        "telefono": E.get("c_tel", ""),
        "email": E.get("c_email", ""),
        "fecha": E.get("c_fecha", date.today()),
        "numero": E.get("cot_numero", ""),
    }


def totales_lineas(lineas: list[dict]) -> dict:
    entradas = [{"cantidad": _numero(ln.get("cantidad"), 0),
                 "precio_neto": _numero(ln.get("unitario"), 0)}
                for ln in lineas if _numero(ln.get("cantidad"), 0) > 0]
    return rv.totales_de(entradas) if entradas else {"neto": 0, "iva": 0, "total": 0}


def mostrar_totales_vivos(totales: dict):
    a, b, c = st.columns(3)
    a.metric("Neto", rv.clp(totales.get("neto", 0)))
    b.metric("IVA (19 %)", rv.clp(totales.get("iva", 0)))
    c.metric("Total con IVA", rv.clp(totales.get("total", 0)))


def _leer_etiquetas(ws):
    for fila in ws.iter_rows():
        etiquetas = {ct.normalizar(c.value): c.column for c in fila if isinstance(c.value, str)}
        if "CANTIDAD" in etiquetas and "DESCRIPCION" in etiquetas:
            return fila[0].row, etiquetas
    raise ValueError("No encontré los encabezados Cantidad y Descripción en el Excel.")


def guardar_xlsx_actual():
    """Actualiza el Excel del borrador con las líneas editadas y lo copia al historial."""
    cot_id = E.get("cot_id")
    if not cot_id or not E.get("lineas"):
        return
    lineas = [ln for ln in E["lineas"] if str(ln.get("descripcion", "")).strip()
              and _numero(ln.get("cantidad"), 0) > 0]
    ruta_actual = Path(E["ruta"]) if E.get("ruta") else None
    origen_pedido = E.get("origen_pedido", "pedido")

    if origen_pedido == "pedido" or not ruta_actual or not ruta_actual.exists():
        tmp = Path(tempfile.mkdtemp()) / "cotizacion.xlsx"
        ct.crear_cotizacion_xlsx(
            [{"cantidad": ln["cantidad"], "descripcion": ln["descripcion"],
              "precio_neto": ln.get("unitario")} for ln in lineas],
            _cliente_actual(), tmp,
        )
        fuente = tmp
    else:
        wb = load_workbook(ruta_actual)
        ws = wb.active
        try:
            fila_enc, columnas = _leer_etiquetas(ws)
        except ValueError:
            wb.close()
            fila_enc, columnas = None, {}
        c_cant, c_desc = columnas.get("CANTIDAD"), columnas.get("DESCRIPCION")
        c_unit, c_total = columnas.get("PUNITARIO"), columnas.get("TOTAL")
        total_rows = {}
        if fila_enc is not None:
            for fila in ws.iter_rows(min_row=fila_enc + 1):
                for celda in fila:
                    if isinstance(celda.value, str):
                        etiqueta = ct.normalizar(celda.value.rstrip(":"))
                        if etiqueta in ("TOTALNETO", "IVA", "TOTAL") and celda.column < ws.max_column:
                            total_rows.setdefault(etiqueta, celda.row)
            inicio = fila_enc + 1
            primera_fila_total = min(total_rows.values()) if total_rows else ws.max_row + 1
            espacios = max(0, primera_fila_total - inicio)
        else:
            espacios = 0

        if fila_enc is None or c_cant is None or c_desc is None or c_unit is None or len(lineas) > espacios:
            if fila_enc is not None:
                wb.close()
            # Caso excepcional: archivo externo sin columnas editables o se agregaron más filas
            # que las disponibles. Se usa la plantilla estándar para no cortar productos.
            cliente = dict(_cliente_actual())
            if ruta_actual and ruta_actual.exists():
                try:
                    d = pc.leer_para_pdf(ruta_actual)
                    pares = {ct.normalizar(e): v for e, v in d["izquierda"] + d["derecha"]}
                    for k, etiqueta in (("nombre", "NOMBRE"), ("rut", "RUT"), ("direccion", "DIRECCION"),
                                        ("ciudad", "CIUDAD"), ("telefono", "TELEFONO"), ("email", "EMAIL")):
                        cliente[k] = cliente.get(k) or pares.get(etiqueta, "")
                except Exception:
                    pass
            tmp = Path(tempfile.mkdtemp()) / "cotizacion.xlsx"
            ct.crear_cotizacion_xlsx(
                [{"cantidad": ln["cantidad"], "descripcion": ln["descripcion"],
                  "precio_neto": ln.get("unitario")} for ln in lineas],
                cliente, tmp,
            )
            fuente = tmp
        else:
            # Las filas de producto ya existentes se reutilizan para conservar la plantilla y su diseño.
            for i in range(espacios):
                fila = inicio + i
                ln = lineas[i] if i < len(lineas) else None
                ws.cell(fila, c_cant).value = ln["cantidad"] if ln else None
                ws.cell(fila, c_desc).value = ln["descripcion"] if ln else None
                ws.cell(fila, c_unit).value = ln.get("unitario") if ln else None
                if c_total:
                    ws.cell(fila, c_total).value = (round(_numero(ln.get("unitario"), 0) * ln["cantidad"])
                                                    if ln and ln.get("unitario") is not None else None)
            tot = totales_lineas(lineas)
            for etiqueta, valor in (("TOTALNETO", tot["neto"]), ("IVA", tot["iva"]), ("TOTAL", tot["total"])):
                fila = total_rows.get(etiqueta)
                if fila:
                    celda_etiqueta = next((c for c in ws[fila]
                                           if isinstance(c.value, str)
                                           and ct.normalizar(c.value.rstrip(":")) == etiqueta), None)
                    if celda_etiqueta and celda_etiqueta.column < ws.max_column:
                        ws.cell(fila, celda_etiqueta.column + 1).value = valor
            wb.save(ruta_actual)
            wb.close()
            actualizar_numero_en_excel(ruta_actual, str(E.get("cot_numero", "")))
            fuente = ruta_actual

    hist.guardar_archivos(int(cot_id), xlsx=fuente)
    registro = hist.obtener(int(cot_id))
    if registro and registro.get("xlsx"):
        E["ruta"] = Path(registro["xlsx"])
        actualizar_numero_en_excel(E["ruta"], str(E.get("cot_numero", "")))
        hist.guardar_archivos(int(cot_id), xlsx=E["ruta"])


def actualizar_inventario(lineas: list[dict]):
    """Aprende precios netos y nombres a medida que se edita la cotización."""
    for ln in lineas:
        if not ln.get("es_producto"):
            continue
        desc = str(ln.get("descripcion", "")).strip()
        clave = str(ln.get("clave_auto", "") or "").strip().upper()
        if not clave and desc:
            clave = ct.clave_descriptiva(desc)
        if clave:
            precio = ln.get("unitario")
            bib.registrar(clave, nombre=desc or clave, alias=desc or None,
                          precio_neto=_numero(precio, 0) if precio is not None and not _es_vacio(precio) and _numero(precio, 0) > 0 else None)


def aplicar_tabla_editada(df: pd.DataFrame):
    """Convierte la tabla visible en líneas de cotización e inventario, conservando filas incompletas."""
    anteriores = E.get("lineas", [])
    lineas = []
    origen_por_etiqueta = {v: k for k, v in ORIGEN.items() if v}
    for i, fila in enumerate(df.to_dict("records")):
        desc = str(fila.get(COL_PROD) or "").strip()
        if desc.lower() == "nan":
            desc = ""
        cantidad = _numero(fila.get(COL_CANT), 0)
        cantidad = max(0, cantidad)
        previo = anteriores[i] if i < len(anteriores) else {}
        codigo_valor = fila.get(COL_COD)
        clave = "" if _es_vacio(codigo_valor) else str(codigo_valor).strip().upper()
        if not clave and desc:
            clave = ct.clave_descriptiva(desc)
        precio_valor = fila.get(COL_PRECIO)
        precio = None if _es_vacio(precio_valor) else max(0, _numero(precio_valor, 0))
        origen_etiqueta = str(fila.get(COL_ORIGEN) or "")
        origen = origen_por_etiqueta.get(origen_etiqueta, previo.get("origen", "descripcion"))
        incluir = fila.get(COL_INCLUIR, bool(desc))
        es_producto = bool(desc) and not any(p in ct.normalizar(desc).lower() for p in ct.PALABRAS_NO_PRODUCTO)
        lineas.append({
            "fila": previo.get("fila", i + 1),
            "cantidad": ct._cantidad(cantidad),
            "descripcion": ct.limpiar_descripcion(desc) if desc else "",
            "unitario": precio,
            "es_producto": es_producto,
            "clave_auto": clave,
            "origen": origen,
            "incluir": not _es_vacio(incluir) and bool(incluir),
        })
    E["lineas"] = lineas
    E["tabla_edit"] = df.copy()
    E["tabla"] = df.copy()
    E.pop("pdf_bytes", None)
    E.pop("pdf_nombre", None)
    E.pop("informe", None)
    actualizar_inventario(lineas)
    guardar_xlsx_actual()
    guardar_estado_borrador()

def editar_lineas_en_etapas():
    """La lista y los precios son editables hasta generar el PDF."""
    if "pdf_bytes" in E:
        st.caption("La cotización quedó bloqueada después de generar el PDF. Inicia otra cotización para hacer cambios.")
        return E.get("tabla_edit")
    base = E.get("tabla_edit", E.get("tabla"))
    if base is None:
        return None
    st.subheader("Productos y precios (editables)")
    st.caption("El precio unitario se ingresa en valor neto, antes de IVA. Puedes editar las celdas o agregar/eliminar filas.")
    key = f"lineas_editor_{E.get('cot_id', 0)}"
    editada = editor(
        base, num_rows="dynamic", hide_index=True, width="stretch", key=key,
        disabled=[COL_ORIGEN],
        column_config={
            COL_INCLUIR: st.column_config.CheckboxColumn(COL_INCLUIR, width="small"),
            COL_CANT: st.column_config.NumberColumn(COL_CANT, min_value=0, step=1, width="small"),
            COL_PROD: st.column_config.TextColumn(COL_PROD, width="large"),
            COL_COD: st.column_config.TextColumn(COL_COD, width="medium"),
            COL_PRECIO: st.column_config.NumberColumn(COL_PRECIO, min_value=0, format="$ %d"),
        },
    )
    if not editada.equals(base):
        aplicar_tabla_editada(editada)
    else:
        E["tabla_edit"] = editada
    mostrar_totales_vivos(totales_lineas(E.get("lineas", [])))
    return editada


def guardar_estado_borrador():
    """Persiste los datos esenciales tras cada interacción de Streamlit."""
    cot_id = E.get("cot_id")
    if not cot_id:
        return
    registro = hist.obtener(int(cot_id))
    if not registro:
        return
    data = registro.get("datos") or {}
    cliente_anterior = data.get("cliente") or {}
    cliente = dict(cliente_anterior)
    for clave, origen in (("nombre", "c_nombre"), ("rut", "c_rut"), ("direccion", "c_dir"),
                          ("ciudad", "c_ciudad"), ("telefono", "c_tel"), ("email", "c_email")):
        if origen in E:
            cliente[clave] = E.get(origen, "")
    fecha = E.get("c_fecha", date.today())
    lineas = E.get("lineas", data.get("lineas", [])) or []
    tot = totales_lineas(lineas) if lineas else {"neto": 0, "iva": 0, "total": 0}
    if not lineas and isinstance(E.get("orden_df"), pd.DataFrame):
        con_iva = E.get("precios_modo", OPCIONES_PRECIO[0]) == OPCIONES_PRECIO[0]
        vistas = []
        for fila in E["orden_df"].to_dict("records"):
            q = _numero(fila.get("Cantidad"), 0)
            p = fila.get("Precio unitario")
            if q > 0 and not _es_vacio(p):
                vistas.append({"cantidad": q, "precio_neto": ct.neto_desde(_numero(p), con_iva)})
        if vistas:
            tot = rv.totales_de(vistas)

    estado = {}
    for clave in ("paso", "vista", "precios_modo", "orden_texto", "c_nombre", "c_rut", "c_dir",
                  "c_ciudad", "c_tel", "c_email", "origen_pedido", "cot_numero", "lineas", "info_desc"):
        if clave in E:
            estado[clave] = _json_seguro(E[clave])
    if "c_fecha" in E:
        estado["c_fecha"] = _json_seguro(E["c_fecha"])
    for clave in ("tabla_edit", "tabla"):
        if isinstance(E.get(clave), pd.DataFrame):
            estado[clave] = _json_seguro(E[clave])
    if isinstance(E.get("tabla_edit"), pd.DataFrame):
        estado["tabla_columns"] = list(E["tabla_edit"].columns)
    if isinstance(E.get("orden_df"), pd.DataFrame):
        estado["orden_df"] = _json_seguro(E["orden_df"])
    if E.get("ruta"):
        estado["ruta"] = str(E["ruta"])
    data.update({"lineas": _json_seguro(lineas), "cliente": _json_seguro(cliente), "_app_state": estado})
    hist.actualizar(int(cot_id), datos=data, fecha=fecha, cliente=cliente.get("nombre", ""),
                    rut=cliente.get("rut", ""), neto=tot["neto"], iva=tot["iva"], total=tot["total"],
                    con_iva=int(E.get("precios_modo", OPCIONES_PRECIO[0]) == OPCIONES_PRECIO[0]),
                    origen=E.get("origen_pedido", registro.get("origen", "pedido")))


def ir_a(paso: int):
    E["paso"] = paso
    st.rerun()


def mostrar_pasos(actual: int):
    nombres = ["Pedido", "Productos", "Fotos y fichas", "PDF"]
    piezas = []
    for i, n in enumerate(nombres, 1):
        clase = "hecho" if i < actual else "activo" if i == actual else ""
        piezas.append(f'<div class="paso {clase}"><b>{"✓" if i < actual else i}</b>{n}</div>')
    st.markdown('<div class="pasos">' + "".join(piezas) + "</div>", unsafe_allow_html=True)


def cargar(ruta: Path, origen_pedido: str):
    """Lee la cotización, identifica cada producto y prepara la tabla de revisión."""
    lineas = ct.leer_cotizacion(ruta)
    productos = [ln for ln in lineas if ln["es_producto"]]
    ia = {}
    if api_key and productos:
        try:
            with st.spinner("Estoy entendiendo cada producto…"):
                resultados = li.entender_productos([ln["descripcion"] for ln in productos], api_key)
            ia = {ln["descripcion"]: r for ln, r in zip(productos, resultados)}
        except Exception:
            st.warning("No pude usar la IA para entender los productos. Sigo con la búsqueda normal.")
    for ln in lineas:
        if ln["es_producto"]:
            ln["clave_auto"], ln["origen"] = ct.resolver_clave(bib, ln["descripcion"], ia.get(ln["descripcion"]))
        else:
            ln["clave_auto"], ln["origen"] = "", ""
    for k in ("pdf_bytes", "pdf_nombre", "informe", "ok_riesgo"):
        E.pop(k, None)
    for ln in lineas:
        ln.setdefault("incluir", bool(ln.get("es_producto")))
    E.update(ruta=ruta, lineas=lineas, info_desc=ia, origen_pedido=origen_pedido,
             tabla=pd.DataFrame({
                 COL_INCLUIR: [ln.get("incluir", ln["es_producto"]) for ln in lineas],
                 COL_CANT: [ln["cantidad"] for ln in lineas],
                 COL_PROD: [ln["descripcion"] for ln in lineas],
                 COL_COD: [ln["clave_auto"] for ln in lineas],
                 COL_ORIGEN: [ORIGEN[ln["origen"]] for ln in lineas],
                 COL_PRECIO: [ln.get("unitario") for ln in lineas]}))
    E["tabla_edit"] = E["tabla"]
    if E.get("cot_id"):
        hist.guardar_archivos(int(E["cot_id"]), xlsx=ruta)
        guardado = hist.obtener(int(E["cot_id"]))
        if guardado and guardado.get("xlsx"):
            E["ruta"] = Path(guardado["xlsx"])


def datos_productos():
    tabla = E.get("tabla_edit")
    if tabla is None or tabla.empty:
        return [], {}
    descripciones, claves = {}, []
    for _, fila in tabla.iterrows():
        inc = fila[COL_INCLUIR]
        # Checkbox puede venir como bool, numpy.bool_ o NaN
        if inc is None or (isinstance(inc, float) and pd.isna(inc)) or not bool(inc):
            continue
        desc = str(fila[COL_PROD] or "").strip()
        if not desc or desc.lower() == "nan":
            continue
        cantidad = _numero(fila.get(COL_CANT), 0)
        if cantidad <= 0:
            continue
        clave = str(fila[COL_COD] if pd.notna(fila[COL_COD]) else "").strip().upper()
        if not clave or clave == "NAN":
            clave = ct.clave_descriptiva(desc)
        if not clave or clave == "NAN":
            continue
        descripciones.setdefault(clave, desc)
        if clave not in claves:
            claves.append(clave)
    return claves, descripciones


# ============================================================== PASO 1: PEDIDO
def aplicar_cliente_guardado():
    """Al elegir un cliente guardado, llena los datos de abajo."""
    elegido = E.get("cli_sel")
    for c in hist.clientes():
        if c["nombre"] == elegido:
            E.update(c_nombre=c["nombre"], c_rut=c["rut"], c_dir=c["direccion"], c_ciudad=c["ciudad"],
                     c_tel=c["telefono"], c_email=c["email"])
            return
    E.update(c_nombre="", c_rut="", c_dir="", c_ciudad="", c_tel="", c_email="")


def pedido_ui():
    st.write("Cuéntame qué productos hay que cotizar. Después te pido los precios y los datos del cliente.")
    con_iva = st.radio("Los precios que voy a escribir…", OPCIONES_PRECIO, horizontal=True,
                       key="precios_modo") == OPCIONES_PRECIO[0]
    FUENTES = ["✍️ Lo escribo o pego aquí", "📊 Subir Excel o CSV", "📷 Subir PDF o foto"]
    fuente = st.radio("¿Cómo me das el pedido?", FUENTES, horizontal=True)
    texto, archivo_orden, archivos_ia = "", None, []
    if fuente == FUENTES[0]:
        texto = st.text_area("Escribe un producto por línea, con su cantidad adelante",
                             placeholder="10 Tóner HP CF283A\n15 Tóner Kyocera TK-3182\n2 Resma carta",
                             height=170, key="orden_texto")
    elif fuente == FUENTES[1]:
        archivo_orden = st.file_uploader("Sube el Excel o CSV del pedido", type=["xlsx", "csv"],
                                         key="orden_archivo")
    else:
        archivos_ia = st.file_uploader("Sube el PDF o las fotos del pedido (puedes subir varias)",
                                       type=["pdf", "jpg", "jpeg", "png", "webp"],
                                       accept_multiple_files=True, key="orden_ia")
        st.caption("El pedido se envía a OpenAI para leerlo. Después revisa que haya quedado bien.")
        if not api_key:
            st.warning("Para leer PDF o fotos necesito la clave de OpenAI (⚙️ Ajustes, arriba a la izquierda).")

    if st.button("📖 Leer el pedido", type="primary", width="stretch"):
        try:
            if fuente == FUENTES[2]:
                if not archivos_ia:
                    raise ValueError("Primero sube un PDF o una foto.")
                with st.spinner("Leyendo el pedido…"):
                    filas = li.leer_orden_ia([(f.name, f.getvalue()) for f in archivos_ia], api_key)
            elif archivo_orden:
                tmp = Path(tempfile.mkdtemp()) / archivo_orden.name
                tmp.write_bytes(archivo_orden.getvalue())
                filas = ct.leer_orden_archivo(tmp)
            else:
                filas = ct.parsear_texto_orden(texto)
        except Exception as e:
            st.error(f"No pude leer el pedido: {e}")
            filas = []
        if not filas:
            st.warning("No encontré productos. Revisa el pedido e inténtalo de nuevo.")
        else:
            precios = []
            for f in filas:  # autocompleta: costo + margen de "Mis productos", o el último precio usado
                clave = bib.detectar(f["descripcion"])
                neto = bib.precio_sugerido(clave) or bib.indice.get(clave, {}).get("precio_neto")
                precios.append(ct.precio_para_mostrar(neto, con_iva) if neto else None)
            E["orden_df"] = pd.DataFrame({"Cantidad": [f["cantidad"] for f in filas],
                                          "Producto": [f["descripcion"] for f in filas],
                                          "Precio unitario": precios})
            E.pop("orden_editor", None)

    if "orden_df" not in E:
        return

    st.divider()
    st.subheader("Productos del pedido")
    st.caption("Revisa las cantidades y los nombres, y escribe el **precio** de cada uno. "
               "Para agregar un producto, escribe en la última fila vacía.")
    orden = editor(
        E["orden_df"], num_rows="dynamic", hide_index=True, width="stretch",
        column_config={"Cantidad": st.column_config.NumberColumn(min_value=0, step=1, width="small"),
                       "Producto": st.column_config.TextColumn("Producto", width="large"),
                       "Precio unitario": st.column_config.NumberColumn(min_value=0, format="$ %d")},
        key="orden_editor")
    E["orden_df"] = orden
    vista_totales = []
    for fila in orden.to_dict("records"):
        cantidad_vista = _numero(fila.get("Cantidad"), 0)
        precio_vista = fila.get("Precio unitario")
        if cantidad_vista > 0 and not _es_vacio(precio_vista):
            vista_totales.append({"cantidad": cantidad_vista,
                                  "precio_neto": ct.neto_desde(_numero(precio_vista), con_iva)})
    st.subheader("Totales de la cotización")
    mostrar_totales_vivos(rv.totales_de(vista_totales) if vista_totales else
                          {"neto": 0, "iva": 0, "total": 0})

    st.subheader("Datos del cliente")
    guardados = hist.clientes()
    if guardados:
        st.selectbox("¿Es un cliente que ya tengo guardado?", ["➕ Cliente nuevo"] + [c["nombre"] for c in guardados],
                     key="cli_sel", on_change=aplicar_cliente_guardado)
    for k in ("c_nombre", "c_rut", "c_dir", "c_ciudad", "c_tel", "c_email"):
        E.setdefault(k, "")
    E.setdefault("c_fecha", date.today())
    c1, c2 = st.columns(2)
    nombre = c1.text_input("Nombre o razón social", key="c_nombre")
    rut = c2.text_input("RUT", key="c_rut")
    if rut and rv.validar_rut(rut) is False:
        c2.caption("⚠️ Este RUT no parece válido. Revisa el dígito final.")
    c3, c4 = st.columns(2)
    direccion = c3.text_input("Dirección", key="c_dir")
    ciudad = c4.text_input("Ciudad", key="c_ciudad")
    c5, c6 = st.columns(2)
    telefono = c5.text_input("Teléfono", key="c_tel")
    email = c6.text_input("Correo (opcional)", key="c_email")
    fecha = st.date_input("Fecha de la cotización", key="c_fecha")

    cot_id = E.get("cot_id")
    registro = hist.obtener(int(cot_id)) if cot_id else None
    numero_actual = E.get("cot_numero") or (registro["numero"] if registro else hist.previsualizar_numero(fecha))
    st.caption(f"🔢 Número de cotización: **{numero_actual}**. Puedes cambiarlo en ⚙️ Ajustes → Número de esta cotización.")

    lineas_orden, sin_precio = [], []
    for _, f in orden.iterrows():
        desc, cant = str(f["Producto"] or "").strip(), f["Cantidad"]
        if not desc or desc == "nan" or pd.isna(cant) or cant <= 0:
            continue
        precio = None if pd.isna(f["Precio unitario"]) or f["Precio unitario"] <= 0 else float(f["Precio unitario"])
        if precio is None:
            sin_precio.append(desc[:40])
        lineas_orden.append({"cantidad": ct._cantidad(float(cant)), "descripcion": ct.limpiar_descripcion(desc),
                             "precio_neto": None if precio is None else ct.neto_desde(precio, con_iva)})
    continuar_igual = True
    if sin_precio:
        st.warning("Todavía faltan precios en: " + ", ".join(sin_precio))
        continuar_igual = st.checkbox("Quiero continuar igual, sin esos precios", key="ok_sin_precio")
    if st.button("Crear la cotización y continuar →", type="primary", width="stretch",
                 disabled=not lineas_orden or not continuar_igual):
        cliente = {"nombre": nombre, "rut": rut, "direccion": direccion, "ciudad": ciudad,
                   "telefono": telefono, "email": email, "fecha": fecha}
        totales = rv.totales_de(lineas_orden)
        datos = (registro.get("datos") or {}) if registro else {}
        datos.update({"lineas": lineas_orden, "cliente": cliente})
        if registro is None:
            cot_id, numero = hist.crear(cliente, lineas_orden, totales, con_iva, fecha,
                                        origen="pedido", numero=numero_actual)
        else:
            numero = numero_actual
            hist.actualizar(int(cot_id), datos=datos, fecha=fecha, cliente=nombre, rut=rut,
                            neto=totales["neto"], iva=totales["iva"], total=totales["total"],
                            con_iva=int(con_iva), origen="pedido")
        E["cot_id"], E["cot_numero"] = int(cot_id), numero
        cliente["numero"] = numero
        base = re.sub(r"[^A-Za-z0-9]+", "_", nombre).strip("_")[:30] or "cliente"
        ruta = Path(tempfile.mkdtemp()) / f"Cotizacion_{re.sub(r'[^A-Za-z0-9._-]', '_', numero)}_{base}.xlsx"
        ct.crear_cotizacion_xlsx(lineas_orden, cliente, ruta)
        hist.guardar_archivos(int(cot_id), xlsx=ruta)
        guardado = hist.obtener(int(cot_id))
        if guardado and guardado.get("xlsx"):
            ruta = Path(guardado["xlsx"])
        hist.guardar_cliente(cliente)
        E["c_fecha"] = fecha
        E["cot_numero"] = numero
        E["num_cot_edit"] = numero
        cargar(ruta, "pedido")
        ir_a(2)


def cotizacion_ui():
    st.write("Si ya tienes la cotización armada en Excel, súbela y yo le agrego las fotos y fichas.")
    archivo = st.file_uploader("Sube la cotización en Excel", type=["xlsx"], key="cot_archivo")
    if archivo and st.button("Continuar →", type="primary", width="stretch", key="cot_continuar"):
        ruta = Path(tempfile.mkdtemp()) / archivo.name
        ruta.write_bytes(archivo.getvalue())
        try:
            cargar(ruta, "cotizacion")
        except ValueError as e:
            st.error(str(e))
            return
        ir_a(2)


def paso_1():
    st.header("1. ¿Qué tienes?")
    if not api_key:
        st.info("💡 Con la clave de OpenAI (en ⚙️ Ajustes, arriba a la izquierda) entiendo mucho mejor los "
                "productos y puedo leer pedidos en PDF o foto.")
    tab_pedido, tab_cot = st.tabs(["📝 Tengo un pedido", "📄 Ya tengo la cotización hecha"])
    with tab_pedido:
        pedido_ui()
    with tab_cot:
        cotizacion_ui()


# ============================================================== PASO 2: PRODUCTOS
def paso_2():
    st.header("2. Revisa tus productos")
    st.write("Esta es la lista. Fíjate en el **código**: con él busco la foto y la ficha técnica. "
             "Si ves uno equivocado o vacío, haz clic en la casilla y corrígelo. Yo lo recordaré para la próxima vez.")
    if not api_key:
        st.info("Sin la clave de OpenAI busco solo con el nombre del producto. Con la clave, deduzco la marca, "
                "el modelo y el código, y la búsqueda funciona mucho mejor.")

    # Usar la última versión editada si existe (evita perder checks al volver)
    editada = editar_lineas_en_etapas()
    if editada is None:
        st.warning("Carga una cotización o lee un pedido en el paso anterior.")
        return

    panel_revision()
    atras, sigue = st.columns(2)
    if atras.button("← Volver", width="stretch"):
        ir_a(1)
    if sigue.button("Siguiente: fotos y fichas →", type="primary", width="stretch"):
        # Forzar que la tabla editada quede guardada y validar ANTES de cambiar de paso
        E["tabla_edit"] = editada
        claves, _ = datos_productos()
        if not claves:
            # Mensaje más claro: ¿falta el check o falta el código?
            marcados = 0
            sin_codigo = 0
            for _, fila in editada.iterrows():
                inc = fila[COL_INCLUIR]
                if inc is not None and not (isinstance(inc, float) and pd.isna(inc)) and bool(inc):
                    marcados += 1
                    clave = str(fila[COL_COD] if pd.notna(fila[COL_COD]) else "").strip()
                    if not clave or clave.upper() == "NAN":
                        sin_codigo += 1
            if marcados == 0:
                st.error("No hay productos seleccionados. Marca al menos uno en la columna **Incluir**.")
            elif sin_codigo:
                st.error(
                    f"Hay {sin_codigo} producto(s) marcado(s) pero **sin código**. "
                    "Escribe un código en la columna Código (o déjalo y se generará solo) e inténtalo de nuevo."
                )
            else:
                st.error("No hay productos seleccionados. Vuelve a marcar al menos uno.")
            return
        ir_a(3)


# ============================================================== PASO 3: FOTOS Y FICHAS
@st.cache_data(show_spinner=False)
def menciones_guardada(ruta: str, mtime: float, clave: str) -> int:
    """Veces que la ficha guardada nombra el código (-1 = no se pudo leer)."""
    try:
        return len(bs.menciones_clave(Path(ruta).read_bytes(), [clave]))
    except Exception:
        return -1


def consultas_sugeridas(clave: str, desc: str) -> tuple[str, str]:
    info = E.get("info_desc", {}).get(desc, {})
    base = bs.limpiar_consulta(desc)
    foto, ficha = info.get("consulta_foto") or base, info.get("consulta_ficha") or base
    if bs.es_codigo(clave):
        foto = foto if ct.normalizar(clave) in ct.normalizar(foto) else f"{clave} {foto}"
        ficha = ficha if ct.normalizar(clave) in ct.normalizar(ficha) else f"{clave} {ficha}"
    return foto, ficha


def detalle_busqueda(d: dict, tipo: str) -> str:
    if not d:
        return ""
    if not d.get("crudos"):
        return "El buscador no respondió (puede ser un límite temporal): intenta de nuevo en un minuto."
    if tipo == "foto":
        return f"Revisé {d['crudos']} imágenes, pero eran muy pequeñas o no se pudieron descargar."
    return f"Revisé {d['crudos']} resultados, pero ninguno era una ficha en PDF utilizable."


def buscar_para(clave: str, q_foto: str, q_ficha: str, fotos=True, fichas=True) -> list[str]:
    avisos = []
    for activo, tipo, funcion, consulta in ((fotos, "foto", bs.buscar_fotos, q_foto),
                                            (fichas, "ficha", bs.buscar_fichas, q_ficha)):
        if activo:
            diag = {}
            try:
                E[f"cand_{tipo}_{clave}"] = funcion(consulta, clave=clave, diag=diag)
            except bs.ErrorBusqueda:
                avisos.append("El buscador no respondió. Intenta de nuevo en un minuto.")
            E[f"diag_{tipo}_{clave}"] = diag
    return avisos


def guardar_candidato(clave: str, tipo: str, c: dict):
    bib.guardar_archivo(clave, tipo, "x" + c["ext"] if tipo == "foto" else "f.pdf", c["datos"])
    E.pop(f"cand_{tipo}_{clave}", None)
    st.rerun()


def tarjeta(clave: str, desc: str):
    foto, ficha = bib.archivos(clave)
    aviso_ficha = ""
    if ficha:
        n = menciones_guardada(str(ficha), ficha.stat().st_mtime, clave) if bs.es_codigo(clave) else -2
        estado_ficha = "⚠️ revisar" if n == 0 else "✅ lista"
        if n == 0:
            aviso_ficha = (f"La ficha guardada no menciona el código «{clave}». Puede ser de otro producto: "
                           "busca otra o revísala.")
    else:
        estado_ficha = "❌ falta"
    with st.container(border=True):
        st.markdown('<span class="tarjeta-marca"></span>', unsafe_allow_html=True)
        c1, c2 = st.columns([1, 3])
        with c1:
            if foto:
                try:
                    st.image(str(foto), width=190)
                except Exception:  # archivo dañado: no debe botar la pantalla
                    st.markdown("🖼️  \n*No pude mostrar la foto. Elige otra.*")
            else:
                st.markdown("🖼️  \n*Sin foto*")
        with c2:
            st.markdown(f"**{desc}**")
            st.caption(f"Código: {clave}")
            st.markdown(f"📷 Foto: {'✅ lista' if foto else '❌ falta'} &nbsp;&nbsp;&nbsp; 📄 Ficha técnica: {estado_ficha}")
            b1, b2 = st.columns(2)
            sug_foto, sug_ficha = consultas_sugeridas(clave, desc)
            q_foto, q_ficha = E.get(f"qf_{clave}", sug_foto), E.get(f"qd_{clave}", sug_ficha)
            if b1.button("🔎 Buscar en internet", key=f"b_{clave}", width="stretch"):
                with st.spinner("Buscando…"):
                    for aviso in buscar_para(clave, q_foto, q_ficha):
                        st.warning(aviso)
            with b2.popover("⚙️ Más opciones", width="stretch"):
                st.text_input("Buscar la foto como", value=sug_foto, key=f"qf_{clave}")
                st.text_input("Buscar la ficha como", value=sug_ficha, key=f"qd_{clave}")
                if ficha and st.button("📄 Sacar la foto desde la ficha", key=f"xf_{clave}"):
                    E[f"cand_foto_{clave}"] = bs.fotos_de_ficha(ficha.read_bytes())
                st.file_uploader("Subir mi propia foto", type=["jpg", "jpeg", "png", "webp"], key=f"foto_{clave}")
                st.file_uploader("Subir mi propia ficha (PDF)", type=["pdf"], key=f"ficha_{clave}")
                url_foto = st.text_input("…o pegar el enlace de una imagen", key=f"urlf_{clave}")
                if st.button("Usar ese enlace de imagen", key=f"dlf_{clave}") and url_foto:
                    try:
                        datos, ext, _, _ = bs.descargar_imagen(url_foto.strip())
                        bib.guardar_archivo(clave, "foto", "x" + ext, datos)
                        st.rerun()
                    except Exception:
                        st.error("No pude usar ese enlace.")
                url_pdf = st.text_input("…o pegar el enlace de un PDF", key=f"urlp_{clave}")
                if st.button("Usar ese enlace de PDF", key=f"dlp_{clave}") and url_pdf:
                    try:
                        datos, _ = bs.descargar_pdf(url_pdf.strip())
                        bib.guardar_archivo(clave, "ficha", "f.pdf", datos)
                        st.rerun()
                    except Exception:
                        st.error("No pude usar ese enlace.")
                st.link_button("Buscar fotos en Google", f"https://www.google.com/search?tbm=isch&q={quote_plus(q_foto)}")
                st.link_button("Buscar ficha en Google",
                               f"https://www.google.com/search?q={quote_plus(q_ficha + ' ficha técnica filetype:pdf')}")
        if aviso_ficha:
            st.warning(aviso_ficha)

        fotos_c = E.get(f"cand_foto_{clave}")
        if fotos_c is not None:
            if not fotos_c:
                st.info("No encontré fotos. " + detalle_busqueda(E.get(f"diag_foto_{clave}", {}), "foto")
                        + " Prueba con «Más opciones».")
            else:
                st.markdown("**Elige la foto que corresponde:**")
                cols = st.columns(3)
                for i, c in enumerate(fotos_c):
                    with cols[i % 3]:
                        st.image(c["datos"], width=260)
                        st.caption(("✅ Coincide con el código" if c["verificada"] else "⚪ Sin confirmar")
                                   + (" · baja calidad" if c.get("baja_resolucion") else ""))
                        if st.button("Elegir esta foto", key=f"uf_{clave}_{i}", width="stretch"):
                            guardar_candidato(clave, "foto", c)
        fichas_c = E.get(f"cand_ficha_{clave}")
        if fichas_c is not None:
            if not fichas_c:
                st.info("No encontré una ficha técnica en PDF. " + detalle_busqueda(
                    E.get(f"diag_ficha_{clave}", {}), "ficha") + " Es normal en productos de librería: "
                        "el PDF llevará solo la foto.")
            else:
                st.markdown("**Elige la ficha técnica que corresponde:**")
                for i, c in enumerate(fichas_c):
                    k1, k2 = st.columns([1, 4])
                    if c["miniatura"]:
                        k1.image(c["miniatura"], width="stretch")
                    if not c.get("verificable", True):
                        estado = "ℹ️ Abre el PDF y confirma que es el producto"
                    else:
                        estado = {3: f"✅ Menciona el código {clave}",
                                  2: f"🟡 Menciona {clave}, pero es largo ({c['paginas']} pág.): puede ser un catálogo",
                                  1: f"⚠️ No menciona {clave}: puede ser de otro producto"}[c["nivel"]]
                    k2.markdown(f"{estado}  \n{c['titulo'][:80]}  \n[Abrir el PDF]({c['url']})")
                    if k2.button("Elegir esta ficha" if c["nivel"] >= 2 else "Elegirla igual",
                                 key=f"ufi_{clave}_{i}"):
                        guardar_candidato(clave, "ficha", c)


def paso_3():
    st.header("3. Fotos y fichas técnicas")
    editar_lineas_en_etapas()
    claves, descripciones = datos_productos()
    if not claves:
        st.warning("No hay productos seleccionados. Vuelve al paso anterior y marca al menos uno.")
        if st.button("← Volver"):
            ir_a(2)
        return
    for clave in claves:  # guarda lo que se haya subido en esta ronda
        for tipo in ("foto", "ficha"):
            subido = E.get(f"{tipo}_{clave}")
            if subido:
                bib.guardar_archivo(clave, tipo, subido.name, subido.getvalue())

    listos = sum(1 for c in claves if all(bib.archivos(c)))
    st.progress(listos / len(claves), text=f"{listos} de {len(claves)} productos tienen foto y ficha")
    faltan = [c for c in claves if not all(bib.archivos(c))]
    if faltan:
        st.write("Presiona el botón y yo busco en internet. Después elige la foto y la ficha correctas en cada producto.")
        if st.button(f"🔎 Buscar fotos y fichas de lo que falta ({len(faltan)})", type="primary", width="stretch"):
            barra, hubo_error = st.progress(0.0, text="Buscando…"), False
            for i, c in enumerate(faltan):
                foto, ficha = bib.archivos(c)
                sug_foto, sug_ficha = consultas_sugeridas(c, descripciones[c])
                hubo_error |= bool(buscar_para(c, E.get(f"qf_{c}", sug_foto), E.get(f"qd_{c}", sug_ficha),
                                               fotos=not foto, fichas=not ficha))
                barra.progress((i + 1) / len(faltan), text=f"Buscando… {i + 1} de {len(faltan)}")
            barra.empty()
            if hubo_error:
                st.warning("El buscador no respondió en algunos productos. Intenta de nuevo en un minuto.")
            else:
                st.success("Listo. Baja y elige la foto y la ficha de cada producto.")
    else:
        st.success("¡Todos los productos tienen foto y ficha!")
    st.caption("Si un producto no tiene ficha técnica, no pasa nada: el PDF llevará solo la foto.")

    for clave in claves:
        tarjeta(clave, descripciones[clave])

    st.divider()
    atras, sigue = st.columns(2)
    if atras.button("← Volver", width="stretch"):
        ir_a(2)
    if sigue.button("Siguiente: crear el PDF →", type="primary", width="stretch"):
        ir_a(4)


# ============================================================== PASO 4: PDF
# ============================================================== REVISIÓN DE LA COTIZACIÓN
def hallazgos() -> list[dict]:
    d = pc.leer_para_pdf(E["ruta"])
    claves = ["" if str(c) == "nan" else str(c or "").strip().upper() for c in E["tabla_edit"][COL_COD]]
    return rv.revisar(d, claves if len(claves) == len(d["lineas"]) else None, bib)


def panel_revision(en_paso_4: bool = False) -> int:
    """Muestra lo que encontró el revisor. Devuelve cuántos errores graves hay."""
    try:
        hs_ = hallazgos()
    except Exception:
        st.info("No pude revisar la cotización automáticamente.")
        return 0
    errores = [h["texto"] for h in hs_ if h["nivel"] == "error"]
    avisos = [h["texto"] for h in hs_ if h["nivel"] == "aviso"]
    if not hs_:
        st.success("✅ Revisé la cotización: los totales, el IVA, los precios y los datos del cliente están en orden.")
        return 0
    st.markdown("#### 🔍 Revisión de la cotización")
    if errores:
        st.error("**Hay que corregir:**\n\n" + "\n".join(f"- {t}" for t in errores))
        if not en_paso_4 and E.get("origen_pedido") == "pedido":
            st.caption("Para corregir precios o datos del cliente, vuelve al paso 1: tu pedido sigue ahí.")
    if avisos:
        st.warning("**Conviene revisar:**\n\n" + "\n".join(f"- {t}" for t in avisos))
    return len(errores)


# ============================================================== PASO 4: PDF
def registrar_en_historial(pdf_bytes: bytes):
    """Guarda el PDF en el historial (y registra la cotización si venía de un Excel ya hecho)."""
    cot_id = E.get("cot_id")
    if cot_id and hist.obtener(cot_id):
        hist.guardar_archivos(cot_id, pdf=pdf_bytes)
        return
    if E.get("origen_pedido") == "cotizacion":
        d = pc.leer_para_pdf(E["ruta"])
        cli = {ct.normalizar(e): v for e, v in d["izquierda"] + d["derecha"]}
        tot = {ct.normalizar(e): v for e, v in d["totales"]}
        numero = hist.numero_libre(str(cli.get("NCOTIZACION") or Path(E["ruta"]).stem))
        lineas = [{"cantidad": ln["cantidad"], "descripcion": ln["descripcion"], "precio_neto": ln["unitario"]}
                  for ln in d["lineas"]]
        cot_id, _ = hist.crear({"nombre": cli.get("NOMBRE", ""), "rut": cli.get("RUT", "")}, lineas,
                               {"neto": tot.get("TOTALNETO") or 0, "iva": tot.get("IVA") or 0,
                                "total": tot.get("TOTAL") or 0},
                               con_iva=False, origen="cotizacion", numero=numero)
        hist.guardar_archivos(cot_id, xlsx=E["ruta"], pdf=pdf_bytes)
        E["cot_id"] = cot_id


def generar_pdf():
    lineas, ruta, tabla = E["lineas"], E["ruta"], E["tabla_edit"]
    finales = []
    for ln, (_, fila) in zip(lineas, tabla.iterrows()):
        if not fila[COL_INCLUIR] or not str(ln.get("descripcion", "")).strip() or _numero(ln.get("cantidad"), 0) <= 0:
            continue
        clave = str(fila[COL_COD]).strip().upper()
        if clave and clave != "NAN":  # aprende: claves corregidas y último precio usado
            aprender = ln.get("origen") != "biblioteca" or ct.normalizar(clave) != ct.normalizar(ln["clave_auto"])
            bib.registrar(clave, nombre=ln["descripcion"], alias=ln["descripcion"] if aprender else None,
                          precio_neto=ln["unitario"] if E.get("origen_pedido") == "pedido" else None)
        else:
            clave = ""
        finales.append({**ln, "clave": clave})
    salida = ruta.with_name(f"{ruta.stem}_con_anexos.pdf")
    with st.spinner("Armando tu PDF…"):
        informe = ct.generar_pdf_final(ruta, finales, bib, salida, motor="libreoffice" if usar_lo else "nativo")
    pdf_bytes = salida.read_bytes()
    registrar_en_historial(pdf_bytes)
    E.update(pdf_bytes=pdf_bytes, pdf_nombre=salida.name, informe=informe)


def paso_4():
    st.header("4. Tu PDF")
    editar_lineas_en_etapas()
    claves, _ = datos_productos()
    con_foto = sum(1 for c in claves if bib.archivos(c)[0])
    con_ficha = sum(1 for c in claves if bib.archivos(c)[1])
    registro = hist.obtener(E["cot_id"]) if E.get("cot_id") else None
    if registro:
        st.markdown(f"**Cotización N° {registro['numero']}** · {registro['cliente'] or 'sin cliente'}")
    st.write(f"La cotización tiene **{len(E['lineas'])} líneas**. Fotos: **{con_foto}** de {len(claves)}. "
             f"Fichas técnicas: **{con_ficha}** de {len(claves)}.")
    n_errores = panel_revision(en_paso_4=True)
    permitido = True
    if n_errores:
        permitido = st.checkbox("Entiendo los problemas de arriba y quiero crear el PDF igual", key="ok_riesgo")
    if st.button("📄 Crear el PDF", type="primary", width="stretch", disabled=not permitido):
        try:
            generar_pdf()
        except Exception as e:
            st.error(f"No pude crear el PDF: {e}")
    if "pdf_bytes" in E:
        st.success(f"¡Listo! Tu PDF tiene {E['informe']['paginas']} páginas y quedó guardado en el historial.")
        st.download_button("⬇️ Descargar el PDF", E["pdf_bytes"], file_name=E["pdf_nombre"],
                           mime="application/pdf", type="primary", width="stretch")
        if E.get("origen_pedido") == "pedido":
            st.download_button("⬇️ Descargar la cotización en Excel", E["ruta"].read_bytes(),
                               file_name=E["ruta"].name, width="stretch")
        textos = (("sin_clave", "Estos productos no tienen código, así que no llevan foto ni ficha"),
                  ("sin_foto", "Les falta la foto"), ("sin_ficha", "Les falta la ficha técnica"),
                  ("ficha_dudosa", "Revisa estas fichas: no mencionan el código del producto"),
                  ("ficha_ilegible", "No pude leer estas fichas"))
        for clave_info, texto in textos:
            if E["informe"][clave_info]:
                st.warning(f"{texto}: " + ", ".join(E["informe"][clave_info]))
    st.divider()
    atras, nueva = st.columns(2)
    if atras.button("← Volver a fotos y fichas", width="stretch"):
        ir_a(3)
    nueva.button("🔄 Hacer otra cotización", width="stretch", on_click=lambda: E.clear())


# ============================================================== HISTORIAL
def fecha_corta(iso: str) -> str:
    try:
        return date.fromisoformat(str(iso)[:10]).strftime("%d/%m/%Y")
    except ValueError:
        return str(iso or "")


def cargar_en_pedido(cot_id: int, mantener: bool):
    """Lleva una cotización guardada al paso 1: como copia nueva, o para seguir editándola."""
    rec = hist.obtener(cot_id)
    if not rec:
        return
    datos = rec["datos"]
    lineas, cli = datos.get("lineas", []), datos.get("cliente", {})
    con_iva = bool(rec["con_iva"])
    for k in ("lineas", "ruta", "tabla", "tabla_edit", "info_desc", "pdf_bytes", "pdf_nombre", "informe",
              "orden_editor", "ok_riesgo", "ok_sin_precio", "cli_sel", "origen_pedido"):
        E.pop(k, None)
    E["precios_modo"] = OPCIONES_PRECIO[0 if con_iva else 1]
    E["orden_df"] = pd.DataFrame({
        "Cantidad": [ln["cantidad"] for ln in lineas],
        "Producto": [ln["descripcion"] for ln in lineas],
        "Precio unitario": [ct.precio_para_mostrar(ln["precio_neto"], con_iva) if ln.get("precio_neto") else None
                            for ln in lineas]})
    E.update(c_nombre=cli.get("nombre", ""), c_rut=cli.get("rut", ""), c_dir=cli.get("direccion", ""),
             c_ciudad=cli.get("ciudad", ""), c_tel=cli.get("telefono", ""), c_email=cli.get("email", ""))
    try:
        E["c_fecha"] = date.fromisoformat(rec["fecha"][:10]) if mantener else date.today()
    except ValueError:
        E["c_fecha"] = date.today()
    if mantener:
        E["cot_id"] = cot_id
    else:
        E.pop("cot_id", None)
    E["paso"], E["vista"] = 1, VISTAS[0]


def historial_cotizaciones():
    E.setdefault("h_ver", 0)
    todos = hist.listar()
    if not todos:
        st.info("Todavía no hay cotizaciones. Cada vez que crees una, quedará guardada aquí con su número.")
        return
    m1, m2, m3 = st.columns(3)
    m1.metric("Cotizaciones", len(todos))
    m2.metric("Monto cotizado", rv.clp(sum(r["total"] or 0 for r in todos)))
    m3.metric("Monto aceptado", rv.clp(sum(r["total"] or 0 for r in todos if r["estado"] == "Aceptada")))
    f1, f2 = st.columns([3, 2])
    texto = f1.text_input("Buscar", placeholder="Número, cliente, RUT o notas", key="h_texto")
    estado = f2.selectbox("Estado", ["Todos"] + hs.ESTADOS, key="h_estado")
    regs = hist.listar(texto, estado)
    if not regs:
        st.warning("No encontré cotizaciones con ese filtro.")
        return
    df = pd.DataFrame({"id": [r["id"] for r in regs], "N°": [r["numero"] for r in regs],
                       "Fecha": [fecha_corta(r["fecha"]) for r in regs], "Cliente": [r["cliente"] for r in regs],
                       "Total": [r["total"] for r in regs], "Estado": [r["estado"] for r in regs],
                       "Notas": [r["notas"] or "" for r in regs]})
    editado = editor(
        df, hide_index=True, width="stretch", key=f"h_tabla_{E['h_ver']}_{texto}_{estado}",
        disabled=["N°", "Fecha", "Cliente", "Total"],
        column_config={"id": None, "Total": st.column_config.NumberColumn(format="$ %d"),
                       "Estado": st.column_config.SelectboxColumn(options=hs.ESTADOS, required=True)})
    if st.button("💾 Guardar estados y notas"):
        cambios = 0
        for (_, fila), reg in zip(editado.iterrows(), regs):
            nuevo_estado, nuevas_notas = fila["Estado"], str(fila["Notas"] or "")
            if nuevo_estado != reg["estado"] or nuevas_notas != (reg["notas"] or ""):
                hist.actualizar(reg["id"], estado=nuevo_estado, notas=nuevas_notas)
                cambios += 1
        E["h_ver"] += 1
        st.success(f"Guardado ({cambios} cambios).")
        st.rerun()

    st.divider()
    opciones = {f"{r['numero']} · {r['cliente'] or 'sin cliente'} · {rv.clp(r['total'] or 0)}": r["id"] for r in regs}
    elegida = st.selectbox("Elige una cotización para ver qué puedes hacer con ella", list(opciones))
    rec = hist.obtener(opciones[elegida])
    with st.container(border=True):
        st.markdown(f"**{rec['numero']}** · {rec['cliente'] or 'sin cliente'} · {rec['estado']} · "
                    f"{fecha_corta(rec['fecha'])}")
        d1, d2 = st.columns(2)
        pdf, xlsx = Path(rec["pdf"] or ""), Path(rec["xlsx"] or "")
        if rec["pdf"] and pdf.exists():
            d1.download_button("⬇️ Descargar el PDF", pdf.read_bytes(), file_name=pdf.name, mime="application/pdf",
                               width="stretch", key=f"dp_{rec['id']}")
        else:
            d1.caption("Todavía no se ha creado el PDF de esta cotización.")
        if rec["xlsx"] and xlsx.exists():
            d2.download_button("⬇️ Descargar el Excel", xlsx.read_bytes(), file_name=xlsx.name,
                               width="stretch", key=f"dx_{rec['id']}")
        b1, b2 = st.columns(2)
        b1.button("🔁 Usar como base de una cotización nueva", width="stretch", key=f"dup_{rec['id']}",
                  on_click=cargar_en_pedido, args=(rec["id"], False))
        b2.button("✏️ Seguir editando esta cotización", width="stretch", key=f"edi_{rec['id']}",
                  on_click=cargar_en_pedido, args=(rec["id"], True))
        with st.expander("🗑️ Eliminar esta cotización"):
            seguro = st.checkbox("Sí, quiero eliminarla para siempre", key=f"del_{rec['id']}")
            if st.button("Eliminar", disabled=not seguro, key=f"delb_{rec['id']}"):
                hist.eliminar(rec["id"])
                E["h_ver"] += 1
                st.rerun()


def clientes_ui():
    E.setdefault("cli_ver", 0)
    cols = ["id", "nombre", "rut", "direccion", "ciudad", "telefono", "email"]
    filas = hist.clientes()
    df = pd.DataFrame(filas, columns=cols) if filas else pd.DataFrame({c: pd.Series(dtype="object") for c in cols})
    st.write("Estos son los clientes que he guardado. Se llenan solos al crear cotizaciones; "
             "también puedes corregirlos o agregar uno en la última fila.")
    editado = editor(
        df, num_rows="dynamic", hide_index=True, width="stretch", key=f"cli_tabla_{E['cli_ver']}",
        column_config={"id": None, "nombre": "Nombre o razón social", "rut": "RUT", "direccion": "Dirección",
                       "ciudad": "Ciudad", "telefono": "Teléfono", "email": "Correo"})
    if st.button("💾 Guardar clientes"):
        hist.reemplazar_clientes(editado.to_dict("records"))
        E["cli_ver"] += 1
        st.success("Clientes guardados.")
        st.rerun()


def vista_historial():
    st.header("📚 Historial")
    tab_cot, tab_cli = st.tabs(["Cotizaciones", "👥 Clientes"])
    with tab_cot:
        historial_cotizaciones()
    with tab_cli:
        clientes_ui()


# ============================================================== MIS PRODUCTOS
def renombrar_producto(clave: str):
    nuevo = str(E.get(f"ren_{clave}", "")).strip().upper()
    try:
        final = bib.renombrar(clave, nuevo, combinar=bool(E.get(f"comb_{clave}")))
        E["p_sel"], E["p_msg"] = final, ("ok", f"Listo: ahora se llama {final}.")
    except ValueError as e:
        E["p_msg"] = ("error", str(e))


def borrar_producto(clave: str):
    bib.eliminar(clave)
    E.pop("p_sel", None)
    E["p_msg"] = ("ok", f"Eliminé {clave}.")


def vista_productos():
    E.setdefault("p_ver", 0)
    st.header("📦 Mis productos")
    st.write("Aquí están los productos que he aprendido. Puedes corregir nombres, anotar el **costo** y el **margen** "
             "para que te sugiera el precio de venta, y cambiar su foto o su ficha.")
    if "p_msg" in E:
        nivel, texto = E.pop("p_msg")
        (st.success if nivel == "ok" else st.error)(texto)
    with st.expander("➕ Agregar un producto a mano"):
        n1, n2 = st.columns(2)
        nuevo_cod = n1.text_input("Código del producto", key="np_cod")
        nuevo_nom = n2.text_input("Nombre", key="np_nom")
        if st.button("Agregar", disabled=not nuevo_cod.strip()):
            bib.registrar(nuevo_cod.strip().upper(), nombre=nuevo_nom.strip() or nuevo_cod.strip())
            E["p_ver"] += 1
            st.rerun()

    filas = bib.filas_catalogo()
    if not filas:
        st.info("Todavía no hay productos guardados. Se irán agregando solos a medida que armes cotizaciones.")
        return
    f1, f2 = st.columns([3, 2])
    texto = f1.text_input("Buscar", placeholder="Código o nombre", key="p_texto")
    solo = f2.checkbox("Solo los incompletos (sin foto o sin ficha)", key="p_solo")
    filas = [f for f in filas if (not texto or texto.lower() in (f["clave"] + " " + f["nombre"]).lower())
             and (not solo or not (f["foto"] and f["ficha"]))]
    if not filas:
        st.warning("No hay productos con ese filtro.")
        return
    mostrar = pd.DataFrame({
        "Código": [f["clave"] for f in filas], "Producto": [f["nombre"] for f in filas],
        "Foto": ["✅" if f["foto"] else "❌" for f in filas], "Ficha": ["✅" if f["ficha"] else "❌" for f in filas],
        "Último precio neto": [f["precio_neto"] for f in filas], "Costo neto": [f["costo_neto"] for f in filas],
        "Margen %": [f["margen_pct"] for f in filas],
        "Sugerido neto": [f["sugerido_neto"] for f in filas],
        "Sugerido con IVA": [None if f["sugerido_neto"] is None else f["sugerido_neto"] * (1 + rv.IVA) for f in filas],
        "Notas": [f["notas"] for f in filas]})
    editado = editor(
        mostrar, hide_index=True, width="stretch", key=f"p_tabla_{E['p_ver']}_{texto}_{solo}",
        disabled=["Código", "Foto", "Ficha", "Último precio neto", "Sugerido neto", "Sugerido con IVA"],
        column_config={c: st.column_config.NumberColumn(format="$ %d") for c in
                       ("Último precio neto", "Costo neto", "Sugerido neto", "Sugerido con IVA")}
                      | {"Margen %": st.column_config.NumberColumn(format="%.1f %%", min_value=0, step=1)})
    st.caption("Precio sugerido = costo × (1 + margen). El margen es sobre el costo: con costo $10.000 y margen 30 % "
               "el precio neto sugerido es $13.000. Al crear un pedido, estos productos aparecen con su precio ya puesto.")
    if st.button("💾 Guardar cambios de la tabla"):
        cambios = 0
        for (_, fila), f in zip(editado.iterrows(), filas):
            nuevos = {"nombre": fila["Producto"], "costo_neto": fila["Costo neto"], "margen_pct": fila["Margen %"],
                      "notas": fila["Notas"] if isinstance(fila["Notas"], str) else ""}
            viejos = {"nombre": f["nombre"], "costo_neto": f["costo_neto"], "margen_pct": f["margen_pct"],
                      "notas": f["notas"] or ""}
            if any(not (pd.isna(nuevos[k]) and viejos[k] is None) and nuevos[k] != viejos[k] for k in nuevos):
                bib.actualizar(f["clave"], **nuevos)
                cambios += 1
        E["p_ver"] += 1
        st.success(f"Guardado ({cambios} productos actualizados).")
        st.rerun()

    st.divider()
    codigos = [f["clave"] for f in filas]
    elegido = st.selectbox("Elige un producto para ver o cambiar su foto y su ficha", codigos, key="p_sel")
    nombre = next(f["nombre"] for f in filas if f["clave"] == elegido)
    tarjeta(elegido, nombre)
    with st.expander("✏️ Cambiar el código o eliminar este producto"):
        st.text_input("Nuevo código", key=f"ren_{elegido}")
        st.checkbox("Si ya existe un producto con ese código, combinarlos", key=f"comb_{elegido}")
        st.button("Cambiar el código", key=f"renb_{elegido}", on_click=renombrar_producto, args=(elegido,))
        st.divider()
        seguro = st.checkbox("Sí, quiero eliminar este producto con su foto y su ficha", key=f"delp_{elegido}")
        st.button("Eliminar producto", disabled=not seguro, key=f"delpb_{elegido}",
                  on_click=borrar_producto, args=(elegido,))


# Recupera borradores de versiones anteriores que tenían Excel guardado pero aún no un estado completo.
if E.get("_reanudar_xlsx"):
    _ruta_reanudar = Path(E.pop("_reanudar_xlsx"))
    _reg_reanudar = hist.obtener(int(E["cot_id"]))
    if _ruta_reanudar.exists() and _reg_reanudar:
        try:
            _cliente_previo = {k: E.get(k, "") for k in
                               ("c_nombre", "c_rut", "c_dir", "c_ciudad", "c_tel", "c_email")}
            cargar(_ruta_reanudar, _reg_reanudar.get("origen", "pedido"))
            E.update(_cliente_previo, paso=2, vista=VISTAS[0],
                     cot_id=int(_reg_reanudar["id"]), cot_numero=_reg_reanudar["numero"],
                     num_cot_edit=_reg_reanudar["numero"])
        except Exception as _error_reanudar:
            st.warning(f"Encontré un borrador anterior, pero no pude reconstruirlo automáticamente: {_error_reanudar}")


# ============================================================== PANTALLA
if TEMA_ACTIVO:
    st.markdown('<h1 class="titulo"><span class="gota"></span>Cotizaciones con fotos y fichas</h1>',
                unsafe_allow_html=True)
else:
    st.title("💧 Cotizaciones con fotos y fichas")
    st.info("🎨 Instalé el nuevo diseño. Presiona la tecla **F5** (actualizar la página) para verlo. "
            "Si sigue igual, cierra la ventana negra y abre el programa de nuevo.")
vista = st.radio("Ir a", VISTAS, horizontal=True, key="vista", label_visibility="collapsed")
if vista == VISTAS[1]:
    vista_historial()
elif vista == VISTAS[2]:
    vista_productos()
else:
    if E["paso"] > 1 and "lineas" not in E:  # sesión nueva sin datos
        E["paso"] = 1
    mostrar_pasos(E["paso"])
    {1: paso_1, 2: paso_2, 3: paso_3, 4: paso_4}[E["paso"]]()

# Guardado automático en SQLite al terminar cada interacción; no depende de pulsar "Siguiente".
try:
    guardar_estado_borrador()
except Exception as _error_guardado:
    st.warning(f"No pude guardar automáticamente este borrador: {_error_guardado}")
